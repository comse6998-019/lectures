"""Section 6 of the notes (The investigation record) as slides, appended to the live deck.

    DECK=../lecture2-1.pptx OUT=/tmp/built.pptx python3 build-section6-slides.py

Reads DECK, appends a Part B divider and twelve content slides, writes OUT. Never deletes a slide,
so no part name is reused. Run `preview.py` on the result before installing it.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from partb_grammar import (GREY, bullets, check, citation, closing_filled, closing_line,
                           code_caption, code_panel, content_slide, divider, notes, panel,
                           table)

DECK = os.environ.get("DECK", "../lecture2-1.pptx")
OUT = os.environ["OUT"]
FIGS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")
MAGENTA = RGBColor(0xC7, 0x35, 0x61)

prs = Presentation(DECK)
start = len(prs.slides)

# ---------------------------------------------------------------- Part B divider
s = divider(prs, "Execution machinery")
notes(s, "Part A asked which decisions we fix at design time. Part B is the machinery that carries "
         "a decision from one step to the next: the record the runtime keeps, the code that reads it "
         "to choose the next node or to stop, and the boundary where a requested action becomes an "
         "executed one.")

# ---------------------------------------------------------------- 1. the record
s = content_slide(prs, "The record every node reads and writes", keep_body=True)
bullets(s, "A graph drawing shows where control may go. One record travels the edge.", [
    [("Declared as a schema. ", 'b'), ("A named set of fields with their types.", '')],
    [("Every node receives it ", 'b'), ("as input and returns an update as output.", '')],
    [("Routing and dispatch ", 'b'), ("are both defined over it, so the rest of Part B reads "
                                      "this record.", '')],
], 700000, 1780000)
panel(prs, s, [("State is ", ''), ("“a shared data structure that represents the current "
                                   "snapshot of your application.”", '')], 2700000)
citation(s, "LangGraph, Graph API", 3341801)
closing_line(s, [("A node returns only the keys it touched.", '')], 3760000)
notes(s, "The picture we drew for each version says where control is allowed to go. It says nothing "
         "about what moves along the edges. What moves is one record, and every mechanism in the rest "
         "of Part B is defined over it. Our controller returns three keys out of the loop's ten.")

# ---------------------------------------------------------------- 2. one transition
s = content_slide(prs, "One transition: propose, then merge")
s.shapes.add_picture(f"{FIGS}/state-transition-deck.png", 250147, 800000, width=6152752)
closing_line(s, [("The node proposes an update; the runtime merges it, one reducer per key.", '')],
             3900000)
notes(s, "Two operations, not one. The node reads the current state together with a result that "
         "arrived from outside the program, and it returns a partial update. The runtime then merges "
         "that update key by key. Each reducer takes two arguments: the value already in state on the "
         "left, the value the node returned on the right. A node can be written without knowing which "
         "of the two its contribution will be.")

# ---------------------------------------------------------------- 3. nondeterminism
s = content_slide(prs, "Why two runs from one state diverge", keep_body=True)
bullets(s, "A transition consumes a result that arrives from outside the program.", [
    [("A model reply ", 'b'), ("is sampled from a distribution.", '')],
    [("A tool observation ", 'b'), ("depends on a fixture the run itself may have modified.", '')],
    [("Neither is determined ", 'b'), ("by the current state alone.", '')],
], 700000, 1560000)
panel(prs, s, [("The same state can produce different successors, so a run can go wrong here even "
                "when every node behaves correctly.", '')], 2500000)
closing_line(s, [("This is the nondeterminism of the task environment, arriving at one point.",
                  '')], 3400000)
notes(s, "This is the term that makes the transition honest. Section 1 called the environment "
         "nondeterministic; here is where that shows up in the code. Same state, same node, different "
         "successor, and nothing has misbehaved.")

# ---------------------------------------------------------------- 4. three records
s = content_slide(prs, "Three records, and who reads each")
table(s,
      [[("Record", '')], [("What it holds", '')], [("Who reads it", '')]],
      [
          [[("Model context", '')],
           [("Instructions, the selected history, and the observations for one call", '')],
           [("The model", '')]],
          [[("Runtime state", '')],
           [("The record kept between steps, including fields that enter no model call", '')],
           [("Nodes, routers, our limits", '')]],
          [[("Durable state", '')],
           [("State written where it survives the end of the process", '')],
           [("A resumed run", '')]],
      ],
      1491615, [1450000, 3064545, 1500000], size=12, height=1680000)
closing_filled(s, [("Model context is assembled from runtime state at ", ''), ("every call", 'b'),
                   (".", '')], 3500000)
notes(s, "These two are easy to confuse because most of their content overlaps. The model sees what "
         "it needs for the current call. The runtime also tracks counters and other fields used to "
         "control execution. Durability: kill the process mid-run and start again, and for the loop as "
         "it stands the agent knows nothing. Checkpointers are a later unit; what the record has to "
         "carry here is the run, not the restart.")

# ---------------------------------------------------------------- 5. the hidden counters
s = content_slide(prs, "The counters the model never sees", keep_body=True)
bullets(s, "Two limits live in fields that enter no model call.", [
    [("Forty model calls. ", 'b'), ("The controller refuses a call once ", ''),
     ("model_calls", 'code'), (" reaches forty.", '')],
    [("Three identical observations. ", 'b'),
     ("The run stops when the last three came back the same.", '')],
    [("Neither number is shown to the model. ", 'b'),
     ("The check is a comparison in the controller.", '')],
], 700000, 1990000)
panel(prs, s, [("A limit the model can read is a limit the model can argue with.", '')], 2870000)
closing_line(s, [("We keep them out on purpose.", '')], 3700000)
notes(s, "The clearest case of a field that stays behind. We want the check to be a comparison in the "
         "controller rather than an instruction in a prompt. Ask the room what happens if you put the "
         "budget in the system prompt instead: which record would you inspect to tell the two designs "
         "apart?")

# ---------------------------------------------------------------- 6. representations
s = content_slide(prs, "How mini-swe-agent represents a state", keep_body=True)
bullets(s, "Three representations of increasing expressive power, all in one agent.", [
    [("Atomic. ", 'b'), ("The run's outcome. Accepted or stopped is a label with no internal "
                         "structure.", '')],
    [("Factored. ", 'b'), ("The record itself. Named fields with values, so two states can agree "
                           "on some and differ on others.", '')],
    [("Structured. ", 'b'), ("The patch and the rejection. A patch is a set of hunks, each over a "
                             "file and a range of lines.", '')],
], 700000, 2010000)
panel(prs, s, [("The expressive representation is the exception, because reasoning over structure "
                "is harder for us and for the model both.", '')], 2890000)
citation(s, "Russell and Norvig, §2.4.7", 3531801)
closing_line(s, [("Keeping the record factored is what buys us reducers, routing and limits.",
                  '')], 3950000)
notes(s, "Atomic, factored, structured is the AIMA axis. The number a benchmark reports is a count of "
         "atomic labels. Structure is confined to the two values where the model has to defend its "
         "work: the patch it proposes and the rejection it gets back.")

# ---------------------------------------------------------------- 7 and 8. the schema, a build pair
SCHEMA_HEAD = [[("Field", '')], [("Reducer", '')], [("What it holds", '')]]
SHIPPED = [
    [[("alert, repo, commit, workspace", 'code')], [("none", '')],
     [("Set once: the alert, the fixture, the pinned commit, the workspace", '')]],
    [[("messages", 'code')], [("append", '')], [("The conversation, in provider form", '')]],
    [[("observations", 'code')], [("append", '')],
     [("One record per dispatched tool call, refused or not", '')]],
    [[("candidate_patch", 'code')], [("replace", '')], [("The patch currently proposed", '')]],
    [[("model_calls, usage", 'code')], [("accumulate", '')],
     [("Calls issued and tokens spent", '')]],
    [[("status", 'code')], [("replace", '')],
     [("Running, or the terminal status that ended the run", '')]],
]
DESIGNED = [
    [[("plan", 'm')], [("replace", '')], [("The subgoals and their status", '')]],
    [[("validation", 'm')], [("replace", '')], [("The latest acceptance decision", '')]],
    [[("reflections", 'm')], [("append", '')], [("Guidance kept from rejected attempts", '')]],
    [[("attempt", 'm')], [("explicit", '')], [("Which attempt is running", '')]],
]
WIDTHS = [1900000, 1050000, 3064545]

s = content_slide(prs, "The state schema, field by field")
table(s, SCHEMA_HEAD, SHIPPED, 1050000, WIDTHS, size=10, height=1920000)
notes(s, "This is what state.py declares today, and what the loop populates. Ten keys. Walk the "
         "reducer column rather than the description: it is the column that decides what a later step "
         "can see.")

s = content_slide(prs, "The state schema, field by field")
table(s, SCHEMA_HEAD, SHIPPED + DESIGNED, 1050000, WIDTHS, size=10,
      rule_before=len(SHIPPED), height=2890000)
closing_filled(s, [("The four below the rule arrive with the planner, the validator and the "
                    "reflection step.", '')], 4140000)
notes(s, "Below the rule are the fields the later versions add. Each arrives with the node that fills "
         "it, so a version's diff shows one change rather than a field that was always waiting. These "
         "four are designed and not yet built, which is what the colour marks throughout Part B.")

# ---------------------------------------------------------------- 9. the schema checks nothing
s = content_slide(prs, "A TypedDict checks nothing at run time")
code_caption(s, "The annotation is for a type checker, and nothing else", 900000)
code_panel(s, [
    [("class AgentState(TypedDict, total=False):", None)],
    [("    model_calls: Annotated[int, operator.add]", None)],
    [("", None)],
    [("# a model reply puts ", GREY), ("\"forty\"", MAGENTA),
     (" here: type-checks, and is wrong at run time", GREY)],
], 1160000, 900000)
panel(prs, s, [("Every guarantee about the contents of state rests on code we write.", '')],
      2350000)
citation(s, "PEP 589", 2991801)
closing_line(s, [("The check that catches it is the argument validation at the tool boundary.",
                  '')], 3420000)
notes(s, "A TypedDict returns an ordinary dictionary at run time. Its types cannot be used in "
         "isinstance tests. So a model reply that puts a string where the schema declares an integer "
         "produces a state that type-checks in the editor and is wrong in the run. Section 8 is where "
         "the check that catches it lives.")

# ---------------------------------------------------------------- 10. reducers
s = content_slide(prs, "Three reducers, three kinds of visibility", keep_body=True)
bullets(s, "A reducer decides what every later step is allowed to see.", [
    [("Append. ", 'b'), ("observations", 'code'), (" and ", ''), ("messages", 'code'),
     (" keep the whole attempt, which is what the stall check compares.", '')],
    [("Replace. ", 'b'), ("candidate_patch", 'code'), (", ", ''), ("plan", 'code'), (" and ", ''),
     ("status", 'code'), (" each name what is true now.", '')],
    [("Accumulate. ", 'b'), ("usage", 'code'), (" sums four token counters and produces no total, "
                                                "so no node ever holds the running total.", '')],
], 700000, 2010000)
panel(prs, s, [("“Each key in the State has its own independent reducer function.”", '')],
      2890000)
citation(s, "LangGraph, Graph API", 3531801)
closing_line(s, [("The four token counters stay apart because they are not interchangeable.",
                  '')], 3950000)
notes(s, "Each of the three answers a question about visibility rather than about storage. The stall "
         "check needs append: it compares the last three observations, which is a question no replaced "
         "field could answer. And a single usage total cannot tell you whether cache reads rose while "
         "fresh input fell.")

# ---------------------------------------------------------------- 11. the reset trap
s = content_slide(prs, "Clearing a field that appends")
code_caption(s, "The reset that does not reset", 800000)
code_panel(s, [
    [("# the reflection step, starting a new attempt", GREY)],
    [("return {\"observations\": ", None), ("[]", MAGENTA), ("}", None)],
    [("# merged: every earlier observation is still exactly where it was", GREY)],
], 1060000, 760000)
panel(prs, s, [("An append reducer given an empty update leaves every earlier observation exactly "
                "where it was.", '')], 2100000)
closing_line(s, [("Clearing an appended field takes an explicit assignment the runtime owns.",
                  '')], 3000000)
notes(s, "This is the one to slow down on. The next attempt is supposed to begin free of the "
         "reasoning that failed, and returning an empty list does not do it. The runtime has to assign "
         "the field explicitly, and it records that assignment in the trace as a state change. "
         "Identity, reflections and the counters survive the reset by design: an attempt limit "
         "enforced on a counter that resets with the attempt is not a limit at all.")

# ---------------------------------------------------------------- 12. the demo
s = content_slide(prs, "One tool result becoming state")
code_caption(s, "Four events in, the dispatcher returns 1,665 characters of source", 750000)
code_panel(s, [
    [("# before: observations [], messages 3, model_calls 1", GREY)],
    [("return {\"observations\": [{\"tool\": \"read_file\", \"ok\": True,", None)],
    [("                          \"result\": \"<1665 chars of source>\"}],", None)],
    [("        \"messages\": [ToolMessage(...)]}", None)],
    [("# merged: observations +1, messages +1, every other key untouched", GREY)],
    [("# the next model input is built from messages alone", GREY)],
], 1010000, 1300000)
panel(prs, s, [("The observation the runtime stored is richer than the message the model receives.",
                '')], 2560000)
closing_line(s, [("State answers what the run knows now; the trace answers what the run did.",
                  '')], 3400000)
notes(s, "From runs/v1-bc284c6b.jsonl, steps 3 and 4. Three things to hold on to. The node named two "
         "keys and the record has ten, so the counters advanced only where a node said so. The stored "
         "observation keeps the arguments and the ok flag alongside the text, and routing and "
         "accounting read those fields rather than the prose. And model_calls stood at one throughout "
         "and was never shown, which is how a budget stays a runtime control. The trace is a fourth "
         "record, and it is what section 7 rests on.")

# ----------------------------------------------------------------
bad = check(prs, start)
if bad:
    print("GEOMETRY FAILURES")
    for b in bad:
        print("  " + b)
    sys.exit(1)
prs.save(OUT)
print(f"{len(prs.slides) - start} slides appended, {start + 1} to {len(prs.slides)} -> {OUT}")

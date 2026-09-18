"""Section 8 of the notes (The tool boundary) as slides.

    DECK=../lecture2-1.pptx OUT=/tmp/built.pptx python3 build-section8-slides.py

Appends twelve content slides. Run after build-section7-slides.py, which this follows in the deck.
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from partb_grammar import (GREY, bullets, check, citation, closing_filled, closing_line,
                           code_caption, code_panel, content_slide, notes, panel, table)

DECK = os.environ.get("DECK", "../lecture2-1.pptx")
OUT = os.environ["OUT"]
FIGS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures")
MAGENTA = RGBColor(0xC7, 0x35, 0x61)

prs = Presentation(DECK)
start = len(prs.slides)

# ---------------------------------------------------------------- 1. the contract
s = content_slide(prs, "What a tool contract contains", keep_body=True)
bullets(s, "A tool has two faces, and only one of them is visible to the model.", [
    [("The model reads ", 'b'), ("a name, a description, and an argument schema.", '')],
    [("The runtime holds ", 'b'), ("the result type, the side-effect class, and the constraints "
                                   "on where the tool may operate.", '')],
    [("Calling the first face the tool ", 'b'), ("is the mistake this section exists to prevent.",
                                                 '')],
], 700000, 1850000)
panel(prs, s, [("A tool contract is the complete specification of one tool, and it has five parts.",
                '')], 2750000)
closing_line(s, [("Two of the five reach the model. Three are properties of the implementation.",
                  '')], 3620000)
notes(s, "Name and purpose; the typed arguments and their constraints; the typed success result and "
         "the error representation; the side-effect class; the execution constraints. The "
         "model-facing definition covers the first two. A model that has read the description knows "
         "nothing about the other three.")

# ---------------------------------------------------------------- 2. the dispatch cycle
s = content_slide(prs, "One request through the dispatcher")
s.shapes.add_picture(f"{FIGS}/dispatch-cycle-deck.png", 250147, 900000, width=6152752)
closing_line(s, [("Every path produces an observation, which is what lets a run continue after a "
                  "refusal.", '')], 3150000)
notes(s, "Five operations, and each of them can end the cycle. The two dashed branches are refusals, "
         "and they rejoin at the same object the successful path produces. What state and the trace do "
         "with that observation is the previous section's material.")

# ---------------------------------------------------------------- 3. the five operations
s = content_slide(prs, "Five operations, each able to stop", keep_body=True)
bullets(s, "Between a model's request and a result, five things happen.", [
    [("Resolve the name. ", 'b'), ("A request naming a tool outside the registry is refused, and "
                                   "the refusal names the tools that do exist.", '')],
    [("Bind the arguments. ", 'b'), ("Binding checks names and arity, so a wrong argument name "
                                     "becomes a readable refusal.", '')],
    [("Check the policy. ", 'b'), ("Every path argument is resolved inside the workspace root "
                                   "before any file is opened.", '')],
    [("Execute, then observe. ", 'b'), ("Success and failure alike become an observation and a ",
                                        ''), ("tool_result", 'code'), (" event.", '')],
], 700000, 2620000)
panel(prs, s, [("The registry is built from the tool enumeration, so there is one list and no way "
                "to advertise a tool without registering it.", '')], 3510000)
notes(s, "Binding checks names and arity and not types, so unambiguous type mismatches are coerced "
         "rather than spending a model call to complain about a string where an integer was wanted.")

# ---------------------------------------------------------------- 4. refusal as observation
s = content_slide(prs, "Why a refusal is an observation", keep_body=True)
bullets(s, "A refused request and a permitted call that fails have the same shape.", [
    [("An exception ends the run. ", 'b'), ("It stays for the genuinely unrecoverable, such as a "
                                            "trace that cannot be written.", '')],
    [("An observation goes back to the model, ", 'b'), ("which reads what went wrong and chooses "
                                                        "again.", '')],
    [("When edit refuses an ambiguous target, ", 'b'), ("the model retries with more context.",
                                                        '')],
], 700000, 2000000)
panel(prs, s, [("A run that raised instead would die on its first imprecise edit.", '')], 2850000)
closing_line(s, [("Only one of the two ever reached the implementation.", '')], 3720000)
notes(s, "This difference is what keeps the reactive loop alive. Both events become observations the "
         "model can read, and the ok flag is what tells them apart.")

# ---------------------------------------------------------------- 5. the signature is the contract
s = content_slide(prs, "Why config is not a parameter")
code_caption(s, "The dispatcher binds the model's arguments against the signature", 800000)
code_panel(s, [
    [("def read_file(ws: Path, path: str, start: int | None = None,", None)],
    [("              end: int | None = None) -> str:", None)],
    [("", None)],
    [("inspect.signature(function).bind(self.ws, **args)", None)],
    [("# add a ", GREY), ("config", MAGENTA), (" parameter and the model can set it", GREY)],
], 1060000, 1250000)
panel(prs, s, [("Every parameter in a tool's signature is a value the model is permitted to "
                "supply.", '')], 2600000)
closing_line(s, [("The workspace root is a parameter, and the dispatcher is what supplies it.",
                  '')], 3450000)
notes(s, "Which is why tools read the run's configuration from the module rather than accepting it as "
         "a parameter. A config parameter would be bindable, and a bindable parameter is one the model "
         "can set. The model never sees the workspace root in the schema and cannot pass it.")

# ---------------------------------------------------------------- 6. six artifacts
s = content_slide(prs, "Six artifacts, six different claims")
table(s,
      [[("Artifact", '')], [("What it contains", '')], [("What it establishes", '')]],
      [
          [[("Model input", '')], [("The tool definitions and the conversation so far", '')],
           [("Which operations were offered, and on what description", '')]],
          [[("Model response", '')], [("A tool call: an identifier, a name, arguments", '')],
           [("That the model requested an operation. Nothing more", '')]],
          [[("tool_request", 'code')], [("The name and arguments as the dispatcher received them",
                                         '')],
           [("That the request reached the dispatcher", '')]],
          [[("Dispatcher", '')], [("Registry lookup, argument binding, path resolution", '')],
           [("That the request satisfied the contract, or why it did not", '')]],
          [[("tool_result", 'code')], [("ok=true, 1,665 characters", '')],
           [("That the function ran and returned", '')]],
          [[("State and next input", '')], [("One observation appended, one tool message appended",
                                             '')],
           [("What the next model call was allowed to see", '')]],
      ],
      1100000, [1500000, 2100000, 2414545], size=10, height=2540000)
closing_filled(s, [("Rows two and five are the pair worth separating.", '')], 3830000)
notes(s, "One read_file call on testcode/BenchmarkTest00283.py, through the artifacts it leaves, from "
         "runs/v1-bc284c6b.jsonl, steps 3 and 4. The provider's raw request stays in the message "
         "history; the trace records the dispatcher's view of it.")

# ---------------------------------------------------------------- 7. what proves a run
s = content_slide(prs, "What proves a function ran", keep_body=True)
bullets(s, "A model response establishes that a tool was requested.", [
    [("A tool call in a response ", 'b'), ("is evidence about the model, and not about the file "
                                           "system.", '')],
    [("A transcript showing a request and a plausible result ", 'b'),
     ("establishes nothing, because the model produced both.", '')],
    [("The ", 'b'), ("tool_result", 'code'), (" event ", 'b'), ("is our record of entering the "
                                                                "function.", '')],
], 700000, 1990000)
panel(prs, s, [("The evidence that a function was entered is the runtime's own record of entering "
                "it.", '')], 2850000)
closing_line(s, [("Several tool calls in one response are dispatched one at a time, each answered.",
                  '')], 3700000)
notes(s, "A tool call left without an answer makes the next model message malformed, which is why the "
         "dispatcher returns one observation for each, including the ones that failed.")

# ---------------------------------------------------------------- 8. the ACI
s = content_slide(prs, "The agent-computer interface", keep_body=True)
bullets(s, "Four principles come out of that work, and each has a counterpart in our dispatcher.", [
    [("Actions are simple and few.", 'b')],
    [("Important operations are consolidated, ", 'b'), ("so that one action makes real progress.",
                                                        '')],
    [("Feedback is informative and concise.", 'b')],
    [("Guardrails catch the mistakes models make, ", 'b'), ("so an error does not propagate "
                                                            "through the turns that follow.", '')],
], 700000, 2390000)
panel(prs, s, [("An agent-computer interface specifies the commands available to the model, and "
                "how environment state is communicated back to it.", '')], 3280000)
citation(s, "Yang et al., §2", 3921801)
notes(s, "SWE-agent is the clearest evidence that the action set and the shape of the feedback move "
         "outcomes, because it holds the model fixed and varies the interface. With GPT-4 Turbo the "
         "resulting system resolves 12.47% of the 2,294 SWE-bench test instances, against 3.8% for the "
         "previous best non-interactive retrieval-augmented system, and the same interface carries to "
         "another model, with Claude 3 Opus resolving 10.5%.")

# ---------------------------------------------------------------- 9. the ablations
s = content_slide(prs, "What each interface component is worth")
table(s,
      [[("Component", '')], [("Variant", '')], [("Resolved", '')], [("Change", '')]],
      [
          [[("SWE-agent", '')], [("as configured", '')], [("18.0", '')], [("", '')]],
          [[("Editor", '')], [("edit action without linting", '')], [("15.0", '')],
           [("−3.0", '')]],
          [[("", '')], [("no edit action at all", '')], [("10.3", '')], [("−7.7", 'b')]],
          [[("Search", '')], [("iterative search", '')], [("12.0", '')], [("−6.0", '')]],
          [[("", '')], [("no search tools", '')], [("15.7", '')], [("−2.3", '')]],
          [[("File viewer", '')], [("30-line window", '')], [("14.3", '')], [("−3.7", '')]],
          [[("", '')], [("whole file", '')], [("12.7", '')], [("−5.3", '')]],
          [[("Context", '')], [("full history", '')], [("15.0", '')], [("−3.0", '')]],
          [[("", '')], [("no demonstration", '')], [("16.3", '')], [("−1.7", '')]],
      ],
      1100000, [1300000, 2300000, 900000, 1514545], size=10, height=2500000,
      rule_before=1)
closing_filled(s, [("Percent of SWE-bench Lite resolved, with the base model ", ''),
                   ("held fixed", 'b'), (".", '')], 3780000)
notes(s, "Table 3 of the SWE-agent paper. Read down the Change column rather than across: each row "
         "varies one component of one system.")

# ---------------------------------------------------------------- 10. what it does not support
s = content_slide(prs, "What the ablation table does not say", keep_body=True)
bullets(s, "Two readings of that table are worth separating.", [
    [("It supports ", 'b'), ("the claim that interface choices move outcomes by several points "
                             "each, and the 7.7-point edit gap is larger than the gap between many "
                             "published systems.", '')],
    [("It does not support ", 'b'), ("a general ranking of interfaces: one system, one base model, "
                                     "a 300-instance subset.", '')],
    [("The row to dwell on ", 'b'), ("is search, where returning results one at a time scores 12.0, "
                                     "below having no search tool at all at 15.7.", '')],
], 700000, 2300000)
panel(prs, s, [("A well-intentioned interface can be worse than nothing when it invites the model "
                "to work through every match in turn.", '')], 3150000)
notes(s, "A variant that helps this agent with this model on these instances may not transfer. That "
         "is the discipline to carry into HW1: an ablation tells you about the system it was run on.")

# ---------------------------------------------------------------- 11. execution authority
s = content_slide(prs, "Each guarantee, and what enforces it")
table(s,
      [[("Guarantee", '')], [("Where it is enforced", '')]],
      [
          [[("Only registered tools run", '')],
           [("The registry, consulted before anything is called", '')]],
          [[("Arguments satisfy the contract", '')],
           [("Signature binding and coercion, before the call", '')]],
          [[("No file outside the workspace is touched", '')],
           [("workspace.resolve, on every path argument", '')]],
          [[("Inspection does not modify the fixture", '')],
           [("The tool implementations, and the pinned commit", '')]],
          [[("An exhausted budget prevents another call", '')],
           [("The controller's check, made before the call", '')]],
          [[("A back edge obeys the attempt limit", '')],
           [("The routing function and the attempt counter", '')]],
      ],
      1150000, [3200000, 2814545], size=11, height=2830000)
closing_filled(s, [("No row is enforced by a sentence in a ", ''), ("prompt", 'b'), (".", '')],
               4180000)
notes(s, "workspace.resolve is where an ordinary-looking shortcut fails. It resolves both the root and "
         "the candidate path fully before comparing them, so .. segments collapse and symlinks are "
         "followed. Comparing the two as strings would look equivalent and would not be, because "
         "/tmp/ws is a prefix of the unrelated directory /tmp/ws-evil.")

# ---------------------------------------------------------------- 12. the demo
s = content_slide(prs, "A refused request, side by side", keep_body=True)
bullets(s, "Two requests, the same dispatcher, and only the arguments differ.", [
    [("The request. ", 'b'), ("read_file", 'code'), (" on ", ''), ("../../etc/passwd", 'code'),
     (", constructed rather than one the model made.", '')],
    [("The check. ", 'b'), ("workspace.resolve", 'code'), (" raises ", ''),
     ("PathEscapeException", 'code'), (", the dispatcher catches it, and no file is opened.", '')],
    [("The evidence. ", 'b'), ("A ", ''), ("tool_request", 'code'), (" event, a failed ", ''),
     ("tool_result", 'code'), (", and no entry recorded in the implementation.", '')],
], 700000, 2100000)
panel(prs, s, [("The model proposes, the runtime executes, and the runtime may refuse.", '')],
      2950000)
closing_line(s, [("Name the component that made each decision, and the artifact that proves it.",
                  '')], 3800000)
notes(s, "Waiting for a model to produce a forbidden request is a poor way to demonstrate a boundary, "
         "so we construct one. Set beside the permitted request, the pair shows the branch: identical "
         "tool, identical dispatcher, and the argument decides whether execution happens. That is the "
         "question HW1 puts to you, on the same fixture and the same tools.")

# ----------------------------------------------------------------
bad = check(prs, start)
if bad:
    print("GEOMETRY FAILURES")
    for b in bad:
        print("  " + b)
    sys.exit(1)
prs.save(OUT)
print(f"{len(prs.slides) - start} slides appended, {start + 1} to {len(prs.slides)} -> {OUT}")

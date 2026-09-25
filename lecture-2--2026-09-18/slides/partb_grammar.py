"""The deck's visual grammar as callable components, current to lectures/CLAUDE.md.

`build-section1-slides.py` predates three rulings and must not be copied: it writes the old navy
(#2A2A63, converted to #28285E on 2026-09-17), sets monospace in Menlo, and puts list items at
16 pt navy where the rule is 14 pt `scheme:tx1`. Everything here is written to the rules as they
stand.

A run is a `(text, style)` pair. Styles: '' plain, 'b' bold, 'm' bold magenta (accent4, for a
forward reference or a part not yet built), 'd' bold #348EDC (verdict lead only), 'code' Maple
Mono NF at the surrounding size.
"""
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

NAVY = RGBColor(0x28, 0x28, 0x5E)
DBLUE = RGBColor(0x34, 0x8E, 0xDC)
GREY = RGBColor(0x8A, 0x8A, 0x9A)
MONO = "Maple Mono NF"

BODY_L, BODY_W = Emu(319251), Emu(6014545)
BODY_R = Emu(6333796)
IMG_L, IMG_W = Emu(250147), Emu(6152752)
TITLE_FOOT = 630050          # the title band ends here
BODY_TOP = 700000            # the default first body object
FLOOR = 4790000              # nothing may reach below this
CLOSE_W = Emu(4445035)       # a filled closing box under a table
CLOSE_L = Emu(6333796 - 4445035)


def sub(el, tag, **attrs):
    e = el.makeelement(qn(tag), {k: str(v) for k, v in attrs.items()})
    el.append(e)
    return e


def _scheme(font, val):
    """Colour a run from the scheme. Goes through python-pptx so the solidFill lands before any
    latin element: an appended fill after <a:latin> is out of schema order and PowerPoint drops
    it on load, which turned section dividers black on 2026-09-25."""
    from pptx.enum.dml import MSO_THEME_COLOR
    names = {'tx1': MSO_THEME_COLOR.TEXT_1, 'tx2': MSO_THEME_COLOR.TEXT_2,
             'bg1': MSO_THEME_COLOR.BACKGROUND_1, 'bg2': MSO_THEME_COLOR.BACKGROUND_2,
             'accent1': MSO_THEME_COLOR.ACCENT_1, 'accent2': MSO_THEME_COLOR.ACCENT_2,
             'accent3': MSO_THEME_COLOR.ACCENT_3, 'accent4': MSO_THEME_COLOR.ACCENT_4,
             'accent5': MSO_THEME_COLOR.ACCENT_5, 'accent6': MSO_THEME_COLOR.ACCENT_6}
    font.color.theme_color = names[val]


def add_runs(p, runs, size, on_navy=False):
    """Lay (text, style) pairs into a paragraph. Emphasis is weight in the run's own colour."""
    for text, style in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        if style == 'code':
            r.font.name = MONO
        r.font.bold = style in ('b', 'm', 'd')
        if style == 'm':
            _scheme(r.font, 'accent4')
        elif style == 'd':
            r.font.color.rgb = DBLUE
        elif on_navy:
            _scheme(r.font, 'tx2')
        else:
            _scheme(r.font, 'tx1')
    return p


def content_slide(prs, title, keep_body=False):
    s = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[1])
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0 or (keep_body and idx == 1):
            continue
        ph._element.getparent().remove(ph._element)
    s.shapes.title.text_frame.text = title      # bare run: the layout supplies 24 pt bold
    return s


def divider(prs, text):
    """1_Custom Layout, one 40 pt box centred on the slide axis, and nothing else."""
    s = prs.slides.add_slide(prs.slide_masters[0].slide_layouts[3])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    w = Emu(5000000)
    tb = s.shapes.add_textbox(Emu(3429000) - Emu(int(w / 2)), Emu(2217807), w, Emu(707886))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.name = '+mj-lt'
    _scheme(r.font, 'tx2')
    return s


def panel(prs, slide, runs, top, height=610801, size=14, one_line=False):
    """The navy statement panel: the one sentence the slide exists to assert."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, BODY_L, Emu(top), BODY_W, Emu(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    sh.line.color.rgb = NAVY
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    inset = Emu(137160 if one_line else 182880)
    tf.margin_left = tf.margin_right = inset
    add_runs(tf.paragraphs[0], runs, size, on_navy=True)
    return sh


def bullets(slide, lead, items, top, height, size=14):
    """The layout's content placeholder: bullet-free navy lead-in, then round navy bullets."""
    body = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 1][0]
    body.left, body.top, body.width, body.height = BODY_L, Emu(top), BODY_W, Emu(height)
    tf = body.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    if lead:
        p = tf.paragraphs[0]
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', '0')
        pPr.set('indent', '0')
        sub(pPr, 'a:buNone')
        r = p.add_run()
        r.text = lead
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = NAVY
        first = False
    for runs in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', '400050')
        pPr.set('indent', '-285750')
        sub(pPr, 'a:spcBef').append(p._p.makeelement(qn('a:spcPts'), {'val': '600'}))
        sub(pPr, 'a:spcAft').append(p._p.makeelement(qn('a:spcPts'), {'val': '600'}))
        bu = sub(pPr, 'a:buClr')
        sub(bu, 'a:srgbClr', val='28285E')
        sub(pPr, 'a:buSzPts', val=str(int(size * 100 * 0.9)))
        sub(pPr, 'a:buFont', typeface='+mn-lt')
        sub(pPr, 'a:buChar', char='●')
        add_runs(p, runs, size)
    return body


def _cell(tc, runs, size, bold_navy=False, header=False, rule_above=None):
    tc.margin_left = tc.margin_right = Emu(95250)
    tc.margin_top = tc.margin_bottom = Emu(47625)
    tc.fill.background()
    tf = tc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    sub(p._p.get_or_add_pPr(), 'a:buNone')
    for text, style in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        if style == 'code':
            r.font.name = MONO
        r.font.bold = header or bold_navy or style in ('b', 'm')
        if style == 'm':
            _scheme(r.font, 'accent4')
        elif bold_navy:
            r.font.color.rgb = NAVY
        else:
            _scheme(r.font, 'tx1')
    tcPr = tc._tc.get_or_add_tcPr()
    for tag, w in (('a:lnT', rule_above), ('a:lnB', 12700)):
        if w is None:
            continue
        ln = tcPr.makeelement(qn(tag), {'w': str(w), 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        sf = sub(ln, 'a:solidFill')
        sub(sf, 'a:schemeClr', val='tx1')
        tcPr.insert(0, ln)


def table(slide, header, rows, top, widths, size=12, rule_before=None, height=None):
    """Column-wide, no fills, no vertical rules; label column bold navy, the rest tx1.

    `rule_before` is the 0-based data-row index where the table's sense changes; it gets the
    1.5 pt rule a \\midrule draws in the notes.
    """
    n_rows, n_cols = len(rows) + 1, len(widths)
    gf = slide.shapes.add_table(n_rows, n_cols, BODY_L, Emu(top), BODY_W,
                                Emu(height or 300000 * n_rows))
    tbl = gf.table
    # Body rows are height 0 so they shrink to their content, which makes the frame's own extent
    # the sum of the declared row heights rather than what the table will really occupy. `height`
    # is the rendered total measured with preview.py, set back on the frame so the geometry check
    # and the previewer both see the box the table actually fills.
    tbl._tbl.tblPr.set('firstRow', '1')
    for i, w in enumerate(widths):
        tbl.columns[i].width = Emu(w)
    tbl.rows[0].height = Emu(300000)
    for j, runs in enumerate(header):
        _cell(tbl.cell(0, j), runs, size, header=True, rule_above=19050)
    for i, row in enumerate(rows):
        tbl.rows[i + 1].height = Emu(0)
        above = 19050 if rule_before is not None and i == rule_before else None
        for j, runs in enumerate(row):
            _cell(tbl.cell(i + 1, j), runs, size, bold_navy=(j == 0), rule_above=above)
    if height:
        gf.height = Emu(height)
    return gf


def closing_line(slide, runs, top, height=430000, size=12):
    """Under a list that already has a panel: outlined, no fill, centred bold navy."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, BODY_L, Emu(top), BODY_W, Emu(height))
    sh.fill.background()
    sh.line.color.rgb = NAVY
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(137160)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for text, style in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = MONO if style == 'code' else r.font.name
        r.font.color.rgb = NAVY
    return sh


def closing_filled(slide, runs, top, height=430000, size=12):
    """Under a table: navy fill, narrower than the column, right-aligned to its edge."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, CLOSE_L, Emu(top), CLOSE_W, Emu(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = NAVY
    sh.line.color.rgb = NAVY
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(137160)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, runs, size, on_navy=True)
    return sh


def citation(slide, text, top):
    """Its own top-anchored box, 10.5 pt italic, right edge flush to the column."""
    tb = slide.shapes.add_textbox(BODY_L, Emu(top), BODY_W, Emu(200000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10.5)
    r.font.italic = True
    _scheme(r.font, 'bg2')
    return tb


def code_caption(slide, text, top):
    tb = slide.shapes.add_textbox(BODY_L, Emu(top), BODY_W, Emu(200000))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return tb


def code_panel(slide, lines, top, height, size=11):
    """Rounded, #F2F2F6, wrap off, every paragraph left-aligned."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, BODY_L, Emu(top), BODY_W, Emu(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF6)
    sh.line.fill.background()
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.04
    except Exception:
        pass
    tf = sh.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(150000)
    tf.margin_top = tf.margin_bottom = Emu(60000)
    for i, spans in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        sub(p._p.get_or_add_pPr(), 'a:buNone')
        for text, colour in spans:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.name = MONO
            r.font.color.rgb = colour or NAVY
            if colour is not None and colour != GREY:
                r.font.bold = True
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def check(prs, first_index):
    """Every rule in CLAUDE.md that a geometry can break. Raises on the first failure."""
    bad = []
    for n, s in enumerate(prs.slides, 1):
        if n <= first_index:
            continue
        boxes = []
        for sh in s.shapes:
            if sh.name.startswith('Title'):
                if (sh.left, sh.top, sh.width, sh.height) != (14251, 19250, 6815475, 610800):
                    bad.append(f'{n}: title placeholder moved')
                continue
            boxes.append((sh.name, sh.left, sh.top, sh.width, sh.height))
            if sh.top + sh.height > FLOOR:
                bad.append(f'{n}: {sh.name} reaches {sh.top + sh.height}, below the {FLOOR} floor')
            if sh.top < TITLE_FOOT:
                bad.append(f'{n}: {sh.name} starts at {sh.top}, inside the title band')
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if (a[2] < b[2] + b[4] and b[2] < a[2] + a[4]
                        and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]):
                    bad.append(f'{n}: {a[0]} overlaps {b[0]}')
    return bad

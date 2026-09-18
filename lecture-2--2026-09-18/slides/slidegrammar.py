#!/usr/bin/env python3
"""Read the visual grammar of a deck, and diff it across edits.

    python3 slidegrammar.py inventory <deck.pptx> [--slides 11-23]   # full visual inventory
    python3 slidegrammar.py summary   <deck.pptx>                    # the conventions, aggregated
    python3 slidegrammar.py snapshot  <deck.pptx> <snap.json>        # record the current state
    python3 slidegrammar.py diff      <old.json> <new.json>          # what changed, in grammar terms

Run `snapshot` after handing a deck over, and `diff` after it comes back edited: every difference
is a correction to the house style, which belongs in CLAUDE.md before the next slides are made.
"""
import json
import sys
from collections import Counter, defaultdict

from pptx import Presentation
from pptx.oxml.ns import qn

EMU_IN = 914400


def _hex(el, tag):
    sub = el.find(qn(tag)) if el is not None else None
    if sub is None:
        return None
    srgb = sub.find(qn('a:srgbClr'))
    if srgb is not None:
        return "#" + srgb.get('val')
    sch = sub.find(qn('a:schemeClr'))
    if sch is not None:
        return "scheme:" + sch.get('val')
    return None


def run_colour(font):
    try:
        if font.color is not None and font.color.type is not None and int(font.color.type) == 1:
            return "#" + str(font.color.rgb)
    except Exception:
        pass
    return _hex(font._rPr, 'a:solidFill') if font._rPr is not None else None


def shape_fill(sh):
    spPr = sh._element.find(qn('p:spPr'))
    if spPr is None:
        return None
    if spPr.find(qn('a:noFill')) is not None:
        return "none"
    return _hex(spPr, 'a:solidFill')


def shape_line(sh):
    spPr = sh._element.find(qn('p:spPr'))
    ln = spPr.find(qn('a:ln')) if spPr is not None else None
    if ln is None:
        return None
    if ln.find(qn('a:noFill')) is not None:
        return {"colour": "none", "pt": None}
    w = ln.get('w')
    return {"colour": _hex(ln, 'a:solidFill'), "pt": round(int(w) / 12700, 2) if w else None}


def para_info(p):
    pPr = p._p.find(qn('a:pPr'))
    d = {"align": str(p.alignment) if p.alignment is not None else None}
    if pPr is not None:
        for a in ('marL', 'indent'):
            if pPr.get(a) is not None:
                d[a] = int(pPr.get(a))
        if pPr.find(qn('a:buNone')) is not None:
            d["bullet"] = None
        bc = pPr.find(qn('a:buChar'))
        if bc is not None:
            d["bullet"] = bc.get('char')
        bclr = pPr.find(qn('a:buClr'))
        if bclr is not None:
            d["bulletColour"] = _hex(pPr, 'a:buClr') or None
        for tag, key in (('a:spcAft', 'spaceAfter'), ('a:spcBef', 'spaceBefore')):
            el = pPr.find(qn(tag))
            if el is not None and el.find(qn('a:spcPts')) is not None:
                d[key] = int(el.find(qn('a:spcPts')).get('val')) / 100
    return d


def shape_info(sh):
    info = {
        "name": sh.name,
        "kind": str(sh.shape_type),
        "left": sh.left, "top": sh.top, "width": sh.width, "height": sh.height,
    }
    if sh.is_placeholder:
        info["placeholder"] = sh.placeholder_format.idx
    if sh.shape_type == 13:
        info["kind"] = "PICTURE"
        return info
    if sh.has_table:
        t = sh.table
        tblPr = t._tbl.find(qn('a:tblPr'))
        sid = tblPr.find(qn('a:tableStyleId')) if tblPr is not None else None
        info["kind"] = "TABLE"
        info["tableStyle"] = sid.text if sid is not None else None
        info["cols"] = [c.width for c in t.columns]
        info["rows"] = [r.height for r in t.rows]
        cell0 = t.cell(0, 0)
        info["cellMargins"] = [cell0.margin_left, cell0.margin_right, cell0.margin_top, cell0.margin_bottom]
        borders = []
        tcPr = cell0._tc.find(qn('a:tcPr'))
        if tcPr is not None:
            for tag in ('a:lnT', 'a:lnB', 'a:lnL', 'a:lnR'):
                el = tcPr.find(qn(tag))
                if el is not None:
                    borders.append({tag[2:]: {"pt": round(int(el.get('w', 0)) / 12700, 2),
                                              "colour": _hex(el, 'a:solidFill')}})
        info["cellBorders"] = borders
        info["text"] = []
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip():
                            info["text"].append({
                                "cell": [ri, ci], "t": r.text,
                                "pt": r.font.size.pt if r.font.size else None,
                                "bold": r.font.bold, "font": r.font.name,
                                "colour": run_colour(r.font)})
        return info
    info["fill"] = shape_fill(sh)
    info["line"] = shape_line(sh)
    if sh.has_text_frame:
        tf = sh.text_frame
        info["anchor"] = str(tf.vertical_anchor) if tf.vertical_anchor is not None else None
        info["margins"] = [tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom]
        info["wrap"] = tf.word_wrap
        info["paras"] = []
        for p in tf.paragraphs:
            runs = [{"t": r.text, "pt": r.font.size.pt if r.font.size else None,
                     "bold": r.font.bold, "italic": r.font.italic,
                     "font": r.font.name, "colour": run_colour(r.font)}
                    for r in p.runs if r.text]
            if runs:
                info["paras"].append({**para_info(p), "runs": runs})
    return info


def inventory(path, only=None):
    prs = Presentation(path)
    out = {"deck": path, "slideW": prs.slide_width, "slideH": prs.slide_height, "slides": []}
    for i, s in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        title = s.shapes.title.text_frame.text.strip() if s.shapes.title is not None else ""
        out["slides"].append({
            "n": i, "layout": s.slide_layout.name, "title": title,
            "shapes": [shape_info(sh) for sh in sorted(s.shapes, key=lambda x: (x.top or 0, x.left or 0))],
        })
    return out


def summary(inv):
    colours, sizes, fonts, lefts, widths, fills, lines, bullets = (Counter() for _ in range(8))
    idioms = Counter()
    for sl in inv["slides"]:
        for sh in sl["shapes"]:
            if sh.get("left") is not None:
                lefts[sh["left"]] += 1
                widths[sh["width"]] += 1
            if sh.get("fill"):
                fills[sh["fill"]] += 1
            if sh.get("line") and sh["line"].get("colour"):
                lines[f'{sh["line"]["colour"]}@{sh["line"]["pt"]}pt'] += 1
            if sh["kind"] == "TABLE":
                idioms["rule-only table" if sh.get("tableStyle", "").startswith("{2D5ABB26") else "table"] += 1
                for t in sh.get("text", []):
                    if t["pt"]:
                        sizes[t["pt"]] += 1
                    if t["colour"]:
                        colours[t["colour"]] += 1
                    if t["font"]:
                        fonts[t["font"]] += 1
            for p in sh.get("paras", []):
                if p.get("bullet"):
                    bullets[p["bullet"]] += 1
                for r in p["runs"]:
                    if r["pt"]:
                        sizes[r["pt"]] += 1
                    if r["colour"]:
                        colours[r["colour"]] += 1
                    if r["font"]:
                        fonts[r["font"]] += 1
            fill, line = sh.get("fill"), (sh.get("line") or {}).get("colour")
            if fill and fill not in ("none", None) and fill != "#FFFFFF":
                idioms[f"filled panel {fill}"] += 1
            elif fill == "none" and line and line != "none":
                idioms[f"outlined box {line}"] += 1
            elif fill == "#FFFFFF" and line and line != "none":
                idioms[f"white box, {line} rule"] += 1
    return {
        "colours": colours.most_common(), "sizes": sizes.most_common(),
        "fonts": fonts.most_common(), "lefts": lefts.most_common(6),
        "widths": widths.most_common(6), "fills": fills.most_common(),
        "lines": lines.most_common(), "bullets": bullets.most_common(),
        "idioms": idioms.most_common(),
    }


def flatten(inv):
    """A comparable key -> value map, so a diff reads as grammar rather than XML."""
    flat = {}
    for sl in inv["slides"]:
        key = f'{sl["n"]:02d} {sl["title"][:40]}'
        for j, sh in enumerate(sl["shapes"]):
            base = f'{key} | {sh["kind"]}#{j} {sh.get("name","")}'
            for f in ("left", "top", "width", "height", "fill", "tableStyle", "wrap", "anchor"):
                if sh.get(f) is not None:
                    flat[f"{base} . {f}"] = sh[f]
            if sh.get("line"):
                flat[f"{base} . line"] = f'{sh["line"]["colour"]}@{sh["line"]["pt"]}pt'
            if sh.get("cols"):
                flat[f"{base} . colWidths"] = sh["cols"]
            if sh.get("rows"):
                flat[f"{base} . rowHeights"] = sh["rows"]
            for t in sh.get("text", []):
                flat[f'{base} . cell{t["cell"]} "{t["t"][:34]}"'] = \
                    f'pt={t["pt"]} bold={t["bold"]} font={t["font"]} colour={t["colour"]}'
            for pi, p in enumerate(sh.get("paras", [])):
                flat[f"{base} . p{pi} props"] = {k: v for k, v in p.items() if k != "runs"}
                for r in p["runs"]:
                    flat[f'{base} . p{pi} "{r["t"][:34]}"'] = \
                        f'pt={r["pt"]} bold={r["bold"]} italic={r["italic"]} font={r["font"]} colour={r["colour"]}'
    return flat


def diff(old, new):
    a, b = flatten(old), flatten(new)
    added = [k for k in b if k not in a]
    removed = [k for k in a if k not in b]
    changed = [(k, a[k], b[k]) for k in a if k in b and a[k] != b[k]]
    return added, removed, changed


def parse_slides(arg):
    out = set()
    for part in arg.split(','):
        if '-' in part:
            lo, hi = part.split('-')
            out.update(range(int(lo), int(hi) + 1))
        else:
            out.add(int(part))
    return out


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    cmd = sys.argv[1]
    if cmd in ("inventory", "summary", "snapshot"):
        only = None
        if "--slides" in sys.argv:
            only = parse_slides(sys.argv[sys.argv.index("--slides") + 1])
        inv = inventory(sys.argv[2], only)
        if cmd == "inventory":
            print(json.dumps(inv, indent=1))
        elif cmd == "summary":
            print(json.dumps(summary(inv), indent=1))
        else:
            with open(sys.argv[3], "w") as fh:
                json.dump(inv, fh, indent=1)
            print(f"snapshot written: {sys.argv[3]} ({len(inv['slides'])} slides)")
    elif cmd == "diff":
        old = json.load(open(sys.argv[2]))
        new = json.load(open(sys.argv[3]))
        added, removed, changed = diff(old, new)
        print(f"=== {len(changed)} changed, {len(added)} added, {len(removed)} removed ===")
        for k, x, y in changed:
            print(f"CHANGED {k}\n    was: {x}\n    now: {y}")
        for k in added:
            print(f"ADDED   {k} = {new and flatten(new)[k]}")
        for k in removed:
            print(f"REMOVED {k}")
    else:
        print(__doc__)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

#!/usr/bin/env python3
"""Lecture 6 (2026-10-16), pass two: the OpenShell guide's figures on the OpenShell slides, and a diagnosis build.

Rahul, 2026-10-15 evening: "Read the openshell docs html and bring that in to the slides." The guide is openshell/openshell-guide.html
(commit be7af99); its figures were redrawn as four-stage builds in figures/r6b/src and rendered to figures/r6b/<name>-<k>.png. This patch
runs on the output of patch-r6.py (pass one, 164 slides) and
  1. puts a figure under the body text of each OpenShell build that can hold one, a stage per click, the closing box under the figure,
  2. adds the four-click build "Diagnose by the boundary that failed." after "Limits of OpenShell enforcement",
  3. renumbers: every slide-jump link text follows its target, and every slide number in the notes that names a slide at or after the
     insertion moves by the four new slides ("the removed slide N" names a slide of the 416-slide lecture 5 deck and stays).
A figure that cannot sit under its build's text (another object holds the band, or it would shrink below MIN_SCALE) is skipped and
logged, and its build keeps Rahul's geometry and text. Rahul's slide text never moves by default. Rahul's notes keep every paragraph
and run as he saved them; the patch only adds paragraphs around them.

usage:  SRC=<pass one> OUT=<pass two> [MEASURES=<json from tools-r3/measure-ppt.py of a previous OUT>] [MOVE=<title>:<k>[,<k>];...] \
        uv run --no-project --python 3.12 --with python-pptx,lxml,pillow python patch-r6b.py
Without MEASURES each figure sits at a first-pass top; with PowerPoint's measurements of the first pass's output, a second run places
each figure 3 pt under the measured body text and each closing box under the figure. Both runs start from SRC, and both take the
same MOVE. MOVE moves a build's paragraphs (0-based) from the slide to the top of its notes so that a skipped figure fits; it is for
a run Rahul has asked for. OFFERED_MOVES lists the two tested offers, with the opening words each paragraph must have:
  MOVE="OpenShell sandboxes an unmodified agent.:1,3"     runtime-startup at 95 percent
  MOVE="OpenShell governs each sandbox on its own.:2"     interceptors at 81 percent
Both together: MOVE="OpenShell sandboxes an unmodified agent.:1,3;OpenShell governs each sandbox on its own.:2". A reference-line
entry whose marker leaves the slide with a moved paragraph goes to the notes with it.
"""
import copy
import importlib.util
import json
import os
import re

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = os.environ["SRC"]
OUT = os.environ["OUT"]
MEASURES = json.load(open(os.environ["MEASURES"])) if os.environ.get("MEASURES") else None
HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures", "r6b")
LOG = []

# patch-r3.py's helpers and constants, imported rather than copied (its module code only reads the environment)
for k in ("L4",):
    os.environ.setdefault(k, SRC)
_spec = importlib.util.spec_from_file_location("patch_r3", os.path.join(HERE, "patch-r3.py"))
r3 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(r3)
title_of, drop, named, set_notes, copy_slide, move_to = r3.title_of, r3.drop, r3.named, r3.set_notes, r3.copy_slide, r3.move_to
is_banner, box_height, NUMREF = r3.is_banner, r3.box_height, r3.NUMREF
GAP, MIN_GAP, FLOOR, TEXT_GAP, EMU_PT = r3.GAP, r3.MIN_GAP, r3.FLOOR, r3.TEXT_GAP, r3.EMU_PT

COL_LEFT, COL_WIDTH = 320040, 6016752
COL_AXIS = COL_LEFT + COL_WIDTH // 2
MIN_SCALE = 0.8           # the brief: a figure shrinks around its top centre, never below 80 percent
DECK_SCALE = 0.9          # lectures/CLAUDE.md: a figure shrinks to 90 percent before the gap under it does
FIRST_TOP = 2743200       # the first pass's figure top, before any measurement
PAD_PX = 7                # 2 pt at 240 dpi: the deck style's border, kept around a figure's drawn extent
NAVY = "28285E"


def log(msg):
    LOG.append(msg)
    print(msg)


# ----------------------------------------------------------------------------------------------- figures
# build title -> (figure, guide chapter, the notes' talking points per stage). The talking points restate the figure's labels, which
# figures/r6b/src/<name>.tikz takes from the guide at be7af99.
FIGURES = {
    "OpenShell sandboxes an unmodified agent.": ("runtime-startup", 7, [
        "The operator's agent image runs unmodified. Once the boundary is provisioned, attached and confirmed and networking is "
        "prepared, the sandbox runtime's launcher thread starts each child.",
        "Each child gives up privilege first: non-root user and group IDs, every capability set dropped, and no_new_privs, a one-way "
        "setting.",
        "The Landlock baseline (ABI v3 or newer, which covers truncation) protects the private /.openshell hierarchy. The user's "
        "filesystem policy is layered above it, so an access must satisfy both.",
        "Seccomp: the network filter sits on the launcher thread and each child inherits it; the child adds its system-call and "
        "self-protection filters, then executes the agent's command, and the supervisor manages the running workload."]),
    "OpenShell policy schema": ("request-enforcement", 9, [
        "Connection admission: the destination's host and port, and the binary's identity.",
        "The request rule on a rest endpoint: method and path; audit logs a violation and lets the request through, enforce blocks it.",
        "Credential binding: only the authorized host, port and path admit the placeholder's substitution.",
        "Middleware may inspect, deny or transform the request before the credential is used; the upstream connection is a separate "
        "verified TLS connection."]),
    "A denied request is recorded for approval.": ("policy-review", 12, [
        "A request is denied, and a narrow rule is drafted, by the supervisor's mapper or through the policy advisor.",
        "The gateway validates the candidate effective policy and the prover reports any new findings; a review token binds this "
        "evaluation to the policy and provider state.",
        "The proposal waits for review, or automatic approval accepts it when there is no new finding and no destination flag; the "
        "approved revision is saved.",
        "The supervisor loads the revision and the workload retries the request. Approval and storage alone do not show that the "
        "revision is loaded."]),
    "Sentry watches from outside the host.": ("network-request", 8, [
        "The workload tool calls connect(); the kernel pauses the call and hands it to the sandbox broker, inside the sandbox.",
        "The supervisor, outside the sandbox, receives the destination with the observed socket and executable evidence, and admits "
        "the TCP connection by host and port, binary and policy generation.",
        "If credentials apply, the supervisor substitutes or signs them, and the connection goes upstream.",
        "Sentry, a BlueField-4 processor on the network path outside the host, watches that traffic and can quarantine the agent; it "
        "never sees a local file read. The figure draws the allowed path only."]),
    "OpenShell governs each sandbox on its own.": ("interceptors", 6, [
        "A create request reaches the gateway, which authenticates it, decodes the permitted operation view, and runs "
        "modify_operation, validate, the normal handler and the post_commit observation.",
        "At creation, an interceptor supplies the approved policy and the provider attachments.",
        "Validation bindings guard every API operation that could weaken the policy later.",
        "Agent identity, agent-to-agent governance and cross-sandbox communication lie outside OpenShell's scope, as Tigera's analysis "
        "puts it."]),
}
STAGES = 4

# Paragraphs that could leave the slide for the notes so that a skipped figure fits (0-based, as Rahul saved them), each with the
# words it must open with. Only background, examples or a paragraph the closing box does not rest on is offered. Nothing moves unless
# MOVE names it: the brief does not authorize moving Rahul's text, so a build whose figure does not fit is skipped and offered.
OFFERED_MOVES = {
    "OpenShell sandboxes an unmodified agent.": {
        1: "NVIDIA announced OpenShell at GTC on March 16, 2026,",      # the announcement history, with the slide's only [2]
        3: "One agent, one sandbox:",                                   # the create command
    },
    "OpenShell governs each sandbox on its own.": {
        2: "Graduated enforcement.",                                    # the AgentKernel comparison, with the link to its slide
    },
}
MOVE_TO_NOTES = {}
for item in filter(None, os.environ.get("MOVE", "").split(";")):
    t, ks = item.rsplit(":", 1)
    MOVE_TO_NOTES.setdefault(t, {}).update({int(k): OFFERED_MOVES.get(t, {}).get(int(k)) for k in ks.split(",")})

# ----------------------------------------------------------------------------------------------- the new build
AFTER = "Limits of OpenShell enforcement"
TEMPLATE = "OpenShell sandboxes an unmodified agent."
NEW_TITLE = "Diagnose by the boundary that failed."
LEAD_IN = "A failed request names an unresolved condition, not a policy that is too strict."
# (bold lead, rest, click that reveals it). The brief's three sentences take ten lines at 12 pt and leave the figure under half its
# width; the slide carries them as one line each and the notes carry them in full.
BULLETS = [
    ("A request can stop", " at any boundary, from its creation to the filesystem.", 1),
    ("Each boundary records", " its own evidence: an error, an event or a result.", 3),
    ("The fix changes", " the narrowest responsible setting, then repeats the request.", 4),
]
NEW_BOX = "Diagnosis starts at the boundary that failed, with that boundary's own evidence. A timeout alone names nothing."
NEW_REF = ("[1] COMS E6998-019, OpenShell: Architecture, Enforcement, and Policy (the course's guide to commit be7af99), chapter 14. "
           "2026.")
NEW_NOTES = ("Chapter 14 of the guide. HW2 asks for the same discipline in the trace: a refused tool result names the place that "
             "refused it (refused_by), so a student reads the boundary, not the symptom. The chapter's table gives one row per "
             "observation; the figure lays the rows along the request path.\n\n"
             "Six boundaries can stop a request: creation, connection admission, TLS and request policy, credential binding, the "
             "upstream service, and the filesystem. Each boundary leaves its own evidence: the create handler's error, the denial "
             "event, the HTTP event and policy mode, the revision result, the path and process identity. The procedure: identify the "
             "exact command and sandbox, read the phase and the effective policy, find the event at that boundary, change the "
             "narrowest responsible setting, repeat the same request.")
NEW_STAGES = [
    "The boundaries in the order a request meets them: creation and provisioning, connection admission, TLS and request policy, "
    "credential binding, the upstream service. The filesystem stands apart.",
    "What the operator sees at each: creation rejected or stuck in Provisioning, connection denied, a connection that opens while "
    "the HTTP operation fails, a saved policy with unchanged behaviour, a failed file write.",
    "The next evidence at each: the RPC error and the create handler; conditions, driver events and supervisor logs; the denial "
    "event and the effective rules; the HTTP event, policy mode and endpoint binding, and the upstream status; the revision result "
    "and a fresh request; the path, process identity and creation policy.",
    "Change the narrowest responsible configuration, then repeat the same request.",
]


# ----------------------------------------------------------------------------------------------- helpers
def build(prs, title):
    idx = [i for i, s in enumerate(prs.slides) if title_of(s) == title]
    assert idx and idx == list(range(idx[0], idx[-1] + 1)), (title, idx)
    return idx


def body_of(slide):
    return next((sh for sh in slide.shapes if sh.name == "Content Placeholder 2"), None)


def title_shape(slide):
    return next(sh for sh in slide.shapes if sh.is_placeholder and sh.placeholder_format.idx == 0)


def banners(slide):
    return sorted([sh for sh in slide.shapes if is_banner(sh)], key=lambda sh: sh.top)


def measured(prs, slide, shape):
    """PowerPoint's record of a shape, read before this run adds the figure. Slide order equals the measured first pass's here; in
    that pass the figure went in right after the body, so a shape after the body sat one place later."""
    if MEASURES is None:
        return None
    shapes = list(slide.shapes)
    k = shapes.index(shape) + 1
    body = body_of(slide)
    if body is not None and k > shapes.index(body) + 1:
        k += 1
    return MEASURES.get(str(list(prs.slides).index(slide) + 1), {}).get(str(k))


def text_bottom(prs, slide, shape):
    m = measured(prs, slide, shape)
    return int(round((m["ttop"] + m["theight"]) * EMU_PT)) if m and m.get("theight") is not None else None


def drop_if_unused(slide, rid):
    if rid and not slide._element.xpath(f"//*[@r:id='{rid}']"):
        slide.part.drop_rel(rid)


def stage_for(k, clicks):
    """Click k of a build shows stage min(k, 4), and its last click the complete figure."""
    return STAGES if k == clicks else min(k, STAGES)


def crop_box(name):
    """The union of the drawn extents of a figure's stages, plus PAD_PX: one rectangle for every stage, so nothing moves. The guide
    style renders an 8 pt white border; cropping it to the deck style's 2 pt makes the gaps the patch sets the gaps a viewer sees."""
    boxes, sizes = [], set()
    for k in range(1, STAGES + 1):
        im = Image.open(os.path.join(FIG, f"{name}-{k}.png")).convert("RGB")
        sizes.add(im.size)
        diff = ImageChops.difference(im, Image.new("RGB", im.size, "white")).convert("L").point(lambda v: 255 if v > 6 else 0)
        boxes.append(diff.getbbox())
    assert len(sizes) == 1, (name, sizes)
    (w, h), = sizes
    l, t = max(0, min(b[0] for b in boxes) - PAD_PX), max(0, min(b[1] for b in boxes) - PAD_PX)
    r, b = min(w, max(b[2] for b in boxes) + PAD_PX), min(h, max(b[3] for b in boxes) + PAD_PX)
    return (w, h), (l, t, r, b)


def stage_notes(lines, clicks):
    out = ["The figure, stage by stage."]
    for k, line in enumerate(lines, 1):
        out.append(f"{k}. {line}")
    if clicks < STAGES:
        rest = " and ".join(str(k) for k in range(clicks, STAGES + 1))
        out.append(f"The last click shows stages {rest} together.")
    return "\n".join(out)


def notes_para(text=None, runs=None):
    """A notes paragraph: plain runs as python-pptx writes them, or (text, attrs) runs; an empty paragraph without either."""
    p = etree.Element(qn("a:p"))
    for t, attrs in (runs if runs is not None else ([(text, {})] if text else [])):
        r = etree.SubElement(p, qn("a:r"))
        if attrs:
            etree.SubElement(r, qn("a:rPr"), attrs)
        etree.SubElement(r, qn("a:t")).text = t
    return p


def body_to_notes(p):
    """A moved body paragraph as a notes paragraph: its words verbatim with their weight and superscript markers; size, colour,
    proofing marks and slide-jump links stay behind. Neighbouring runs of one style merge, so "(slide 147)" is one run and the
    renumbering finds it."""
    runs = []
    for el in p:
        if el.tag == qn("a:r"):
            rpr = el.find(qn("a:rPr"))
            attrs = {}
            if rpr is not None and rpr.get("b") == "1":
                attrs["b"] = "1"
            if rpr is not None and rpr.get("baseline"):
                attrs["baseline"] = rpr.get("baseline")
            text = el.find(qn("a:t")).text or ""
        elif el.tag == qn("a:br"):
            attrs, text = {}, " "
        else:
            assert el.tag in (qn("a:pPr"), qn("a:endParaRPr")), el.tag
            continue
        if runs and runs[-1][1] == attrs:
            runs[-1] = (runs[-1][0] + text, attrs)
        else:
            runs.append((text, attrs))
    return notes_para(runs=runs)


def extend_notes(slide, before, after):
    """Rahul's notes keep every paragraph and run as he saved them (python-pptx's text setter would flatten his bold runs). The
    paragraphs `before` open the notes and `after` close them, each block set off by an empty paragraph, as set_notes("\n\n") did."""
    txBody = slide.notes_slide.notes_text_frame._txBody
    old = txBody.findall(qn("a:p"))
    blank = lambda p: not "".join(t.text or "" for t in p.iter(qn("a:t"))).strip()
    while old and blank(old[-1]):           # trailing empty paragraphs, so the blocks join with one empty paragraph
        txBody.remove(old.pop())
    blocks = [[p] for p in before] + ([old] if old else []) + ([after] if after else [])
    for p in list(txBody.findall(qn("a:p"))):
        txBody.remove(p)                    # detached and re-appended in order; Rahul's elements are the same objects, untouched
    for i, block in enumerate(blocks):
        if i:
            txBody.append(notes_para())
        for p in block:
            txBody.append(p)


CITE_ENTRY = re.compile(r"\[(\d+)\] ")


def trim_citations(slide):
    """After paragraphs left the slide: a reference-line entry whose marker no longer appears on the slide leaves the line, and its
    text is returned for the notes. Only trailing entries can leave, so the markers that stay keep their numbers."""
    cit = named(slide, "Citations")
    if not cit:
        return []
    marks = set()
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name != "Citations":
            marks |= set(re.findall(r"\[(\d+)\]", sh.text_frame.text))
    p = cit[0].text_frame.paragraphs[0]._p
    ts = [r.find(qn("a:t")) for r in p.findall(qn("a:r"))]
    text = "".join(t.text or "" for t in ts)
    entries = [(int(m.group(1)), m.start()) for m in CITE_ENTRY.finditer(text)]
    gone = [n for n, _ in entries if str(n) not in marks]
    if not gone:
        return []
    assert gone == [n for n, _ in entries][-len(gone):], ("only trailing entries may leave", gone, entries)
    cut = entries[-len(gone)][1]
    start = len(text[:cut].rstrip())
    moved = [text[s:(entries[i + 1][1] if i + 1 < len(entries) else len(text))].strip()
             for i, (n, s) in enumerate(entries) if n in gone]
    pos = 0
    for t in ts:                            # cut the text from `start` to the end across the runs, keeping each run's formatting
        s = t.text or ""
        keep = max(0, start - pos)
        if keep < len(s):
            t.text = s[:keep]
        pos += len(s)
    for r in p.findall(qn("a:r")):
        if not (r.find(qn("a:t")).text or ""):
            p.remove(r)
    return moved


def whiten(p):
    """Hidden until its click, as the r2 builder does it: transparent white runs and no bullet glyph; the paragraph keeps its place."""
    for el in list(p.iter(qn("a:rPr"))) + list(p.iter(qn("a:endParaRPr"))):
        for f in el.findall(qn("a:solidFill")):
            el.remove(f)
        fill = etree.Element(qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", "FFFFFF")
        etree.SubElement(clr, qn("a:alpha")).set("val", "0")
        el.insert(0, fill)
    ppr = p.find(qn("a:pPr"))
    if ppr is not None and ppr.find(qn("a:buChar")) is not None:
        ppr.replace(ppr.find(qn("a:buChar")), etree.Element(qn("a:buNone")))


# ----------------------------------------------------------------------------------------------- 1. the new build
def run(text, sz, bold, color, baseline=None):
    r = etree.Element(qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", str(sz))
    rpr.set("b", "1" if bold else "0")
    if baseline:
        rpr.set("baseline", baseline)
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    etree.SubElement(r, qn("a:t")).text = text
    return r


def new_build(prs):
    """Four slides cloned from the last slide of the template build, after the build AFTER. Returns them."""
    tpl = prs.slides[build(prs, TEMPLATE)[-1]]
    tbody = body_of(tpl)
    bullet_ppr = tbody.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert bullet_ppr.find(qn("a:buChar")) is not None
    slides = []
    for k in range(1, STAGES + 1):
        s = copy_slide(prs, tpl)
        t = title_shape(s)
        rs = t.text_frame.paragraphs[0].runs
        assert len(rs) == 1 and rs[0].text == TEMPLATE, [r.text for r in rs]
        rs[0].text = NEW_TITLE
        txBody = body_of(s).text_frame._txBody
        for p in txBody.findall(qn("a:p")):
            txBody.remove(p)
        lead = etree.SubElement(txBody, qn("a:p"))
        ppr = etree.SubElement(lead, qn("a:pPr"))
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        etree.SubElement(etree.SubElement(ppr, qn("a:spcAft")), qn("a:spcPts")).set("val", "300")
        etree.SubElement(ppr, qn("a:buNone"))
        lead.append(run(LEAD_IN[:-1], 1400, True, NAVY))
        lead.append(run("[1]", 1400, True, NAVY, baseline="30000"))
        lead.append(run(".", 1400, True, NAVY))
        for bold, rest, click in BULLETS:
            p = etree.SubElement(txBody, qn("a:p"))
            p.append(copy.deepcopy(bullet_ppr))
            p.append(run(bold, 1200, True, "000000"))
            p.append(run(rest, 1200, False, "000000"))
            etree.SubElement(p, qn("a:endParaRPr")).set("lang", "en-US")
            if click > k:
                whiten(p)
        bs = banners(s)
        assert len(bs) == 1, [b.text_frame.text for b in bs]
        if k < STAGES:
            drop(bs[0])
        else:
            r3.set_first_run(bs[0], NEW_BOX)
        cit = named(s, "Citations")
        assert len(cit) == 1
        r3.set_first_run(cit[0], NEW_REF)
        set_notes(s, NEW_NOTES + "\n\n" + stage_notes(NEW_STAGES, STAGES))
        slides.append(s)
    at = build(prs, AFTER)[-1] + 1
    move_to(prs, slides, at)
    log(f"new build: {NEW_TITLE!r} inserted at {at + 1}-{at + STAGES}, after {AFTER!r}; bullets on clicks "
        f"{[c for _, _, c in BULLETS]}, closing box on click {STAGES}")
    return slides, at


# ----------------------------------------------------------------------------------------------- 2. figures under the text
def middle_band(slide):
    """Objects other than the title, the body, closing boxes and the reference line: a code window, a figure, a table."""
    return [sh for sh in slide.shapes if sh.name not in ("Content Placeholder 2", "Citations", "Back link", "Skip link")
            and not (sh.is_placeholder and sh.placeholder_format.idx == 0) and not is_banner(sh)]


def place(prs, slides, name, notes_lines, moved=None, new=False):
    """Stage k of figure `name` on click k of the build `slides`, under the body text; the closing box under the figure."""
    title = title_of(slides[0])
    last = slides[-1]
    band = middle_band(last)
    if band:
        log(f"figures: SKIP {title!r}: {[sh.name for sh in band]} hold the band under the text; {name} is not placed")
        return None
    (wpx, hpx), (l, t, r, b) = crop_box(name)
    w = COL_WIDTH
    h = int(round(COL_WIDTH * (b - t) / (r - l)))
    bs = banners(last)
    assert len(bs) == 1, (title, [x.text_frame.text for x in bs])
    box = bs[0]
    bh = box.height
    if new and MEASURES:
        m = measured(prs, last, box)
        if m and m.get("lines"):
            bh = box_height(int(m["lines"]))
    tb = text_bottom(prs, last, body_of(last))
    top = FIRST_TOP if tb is None else tb + TEXT_GAP
    # lectures/CLAUDE.md: the figure shrinks to 90 percent before the gap shrinks to MIN_GAP; the brief then lets it go to 80 percent
    scale, gap = 1.0, GAP
    if top + h + GAP + bh > FLOOR:
        scale = max(DECK_SCALE, (FLOOR - bh - GAP - top) / h)
        gap = min(GAP, FLOOR - bh - top - int(round(h * scale)))
        if gap < MIN_GAP:
            scale = max(MIN_SCALE, (FLOOR - bh - MIN_GAP - top) / h)
            gap = min(GAP, FLOOR - bh - top - int(round(h * scale)))
    if tb is not None and gap < MIN_GAP:
        need = top + int(round(h * MIN_SCALE)) + MIN_GAP + bh - FLOOR
        offer = OFFERED_MOVES.get(title) if not moved else None
        log(f"figures: SKIP {title!r}: {name} at {MIN_SCALE:.0%} would push the closing box {need / EMU_PT:.1f} pt below the floor "
            f"(text ends at {tb}, box {bh} tall); the build keeps its geometry and text"
            + (f"; offered: MOVE=\"{title}:{','.join(str(k) for k in sorted(offer))}\"" if offer else ""))
        return None
    # the figure fits: the text moves, then the pictures go in
    refs = []
    if moved:
        for s in slides:
            ps = body_of(s).text_frame._txBody.findall(qn("a:p"))
            for k in sorted(moved, reverse=True):
                p = ps[k]
                text = "".join(x.text or "" for x in p.iter(qn("a:t")))
                assert moved[k] is None or text.startswith(moved[k]), (title, k, text[:60])
                rids = [h.get(qn("r:id")) for h in p.iter(qn("a:hlinkClick"))]
                p.getparent().remove(p)
                for rid in rids:
                    drop_if_unused(s, rid)
            refs = trim_citations(s)
        log(f"figures: {title!r}: paragraphs {sorted(moved)} moved to the notes (MOVE)"
            + (f"; reference entries leave the slide with them: {[r[:40] for r in refs]}" if refs else ""))
    fw, fh = int(round(w * scale)), int(round(h * scale))
    left = COL_LEFT if scale == 1.0 else COL_AXIS - fw // 2
    box_top = top + fh + gap
    stages = []
    for k, s in enumerate(slides, 1):
        st = stage_for(k, len(slides))
        stages.append(st)
        pic = s.shapes.add_picture(os.path.join(FIG, f"{name}-{st}.png"), left, top, fw, fh)
        pic.crop_left, pic.crop_top = l / wpx, t / hpx
        pic.crop_right, pic.crop_bottom = (wpx - r) / wpx, (hpx - b) / hpx
        pic.name = f"Figure {name}"
        s.shapes._spTree.remove(pic._element)
        body_of(s)._element.addnext(pic._element)
        for x in banners(s):
            x.top, x.height = box_top, bh
        bk = body_of(s)
        if bk.top + bk.height > top - TEXT_GAP:
            bk.height = top - TEXT_GAP - bk.top
    first = list(prs.slides).index(slides[0]) + 1
    where = f"top {top}" + (f" (3 pt under the measured text)" if tb is not None else " (first pass)")
    log(f"figures: {title!r}: slides {first}-{first + len(slides) - 1} show {name} stages {stages}, {where}, "
        f"{scale:.0%} ({fw} x {fh}, crop {l},{t},{r},{b} of {wpx}x{hpx} px), closing box at {box_top} ({gap / EMU_PT:.1f} pt under)")
    return refs


def figures(prs, new_slides):
    placed = []
    for title, (name, chapter, lines) in FIGURES.items():
        slides = [prs.slides[i] for i in build(prs, title)]
        moved = MOVE_TO_NOTES.get(title)
        # the moved paragraphs, read before any change (a build carries the whole text on every click)
        moved_ps = [] if not moved else [copy.deepcopy(body_of(slides[-1]).text_frame._txBody.findall(qn("a:p"))[k])
                                         for k in sorted(moved)]
        refs = place(prs, slides, name, lines, moved)
        if refs is None:
            continue
        for s in slides:
            before = [body_to_notes(p) for p in moved_ps] + [notes_para(r) for r in refs]
            after = [notes_para(x) for x in stage_notes(lines, len(slides)).split("\n")]
            extend_notes(s, before, after)
        placed.append(name)
    assert place(prs, new_slides, "diagnosis", NEW_STAGES, new=True) is not None, "the new build must hold its figure"
    placed.append("diagnosis")
    log(f"figures placed: {placed}")


# ----------------------------------------------------------------------------------------------- 3. slide numbers
def renumber(prs, at, n_new, new_ids):
    """Link texts follow their targets; in the notes, a number that names a slide at or after the insertion moves by n_new.
    `at` is the 1-based position of the first new slide."""
    index_of = {s.part: i + 1 for i, s in enumerate(prs.slides)}
    linked = shifted = 0
    for s in prs.slides:
        for h in s._element.iter(qn("a:hlinkClick")):
            if h.get("action", "").startswith("ppaction://hlinksldjump"):
                t = h.getparent().getparent().find(qn("a:t"))
                if t is not None and t.text and t.text.strip().isdigit():
                    new = str(index_of[s.part.rels[h.get(qn("r:id"))].target_part])
                    if new != t.text:
                        t.text = new
                        linked += 1
        if s.slide_id in new_ids or not s.has_notes_slide:
            continue
        for t in s.notes_slide.notes_text_frame._txBody.iter(qn("a:t")):
            if not t.text or not NUMREF.search(t.text):
                continue

            def one(m, text=t.text):
                if text[max(0, m.start() - 8):m.start()].lower().endswith("removed "):
                    return m.group(0)       # a slide of the 416-slide deck that this deck no longer carries
                word, sp, first, rest = m.groups()
                bump = lambda n: str(int(n) + n_new) if int(n) >= at else n
                return f"{word}{sp}{bump(first)}" + re.sub(r"\d+", lambda k: bump(k.group(0)), rest)
            new = NUMREF.sub(one, t.text)
            if new != t.text:
                t.text = new
                shifted += 1
    log(f"renumbered: {linked} link texts, {shifted} note runs (numbers from {at} on move by {n_new})")


# ----------------------------------------------------------------------------------------------- main
def main():
    prs = Presentation(SRC)
    n0 = len(prs.slides)
    new_slides, at0 = new_build(prs)
    figures(prs, new_slides)
    renumber(prs, at0 + 1, len(new_slides), {s.slide_id for s in new_slides})
    prs.save(OUT)
    log(f"saved {OUT}: {n0} -> {len(prs.slides)} slides")
    with open(OUT + ".log", "w") as f:
        f.write("\n".join(LOG) + "\n")


if __name__ == "__main__":
    main()

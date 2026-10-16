#!/usr/bin/env python3
"""Lecture 6 (2026-10-16), Task 4.6 step 3: the OpenShell segment rebuilt from first principles, spliced into the installed deck.

Rahul, 2026-10-15 22:20, on the installed pass one: the OpenShell section was "haphazard and not teaching ready". He approved a rebuild
of the whole segment as a first-principles sequence, every build with its own figure, replacing the old builds in place. The new builds
are the fragment deck that openshell/build-openshell.py makes from openshell/slides.openshell.json (a title slide, then builds 1 to 10
and 12 to 15 of the brief). This patch
  0. applies patch-r6.py MODE=fix1 first when the input still lacks it (the midterm scope as a six-click build, the Limits pair, the
     early-poll notes: Task 4.3's fix round, committed in a36b3bc but never installed); a deck that already has the six midterm clicks
     is spliced as it is,
  1. finds the old segment, from the first slide titled "OpenShell architecture and responsibilities" to the last titled "Limits of
     OpenShell enforcement", and in it the build "Real credentials never enter the sandbox.", which stays (build 11 of the brief),
  2. copies every fragment slide but its title slide to the deck (patch-r3.py's copy_slide, the fragment's slide-jump links stripped
     first and re-pointed afterwards: build 13's "skip ›" opens the first click of build 14), with their notes; moves the kept
     credentials build to follow build 10; moves the new run to the old segment's first position; drops the other old slides,
  3. renumbers: every slide-jump link text follows its target; every "slide N" in the notes names N's new position (a slide of the
     old segment names the new build that took over its subject, a range ending at the old segment's end ends at the new segment's
     end); "the removed slide N (...)" names a lecture 5 slide and stays; the copied slides' notes carry no numbers,
  4. drops the "‹ back" link from builds that lost their last incoming slide-number link (it belonged to a link from the old segment),
  5. saves and prints the old range, the new range and the slide count.

usage:  SRC=<snapshot of lecture6.pptx> FRAG=/tmp/l6-openshell-fragment.pptx OUT=<patched deck> [FIX1=0] \
        uv run --no-project --python 3.12 --with python-pptx,lxml python patch-r6c.py
FRAG's .map.json (written by the builder) says which fragment slides hold each build. FIX1=0 skips step 0.
"""
import copy
import importlib.util
import json
import os
import re
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

SRC, FRAG, OUT = os.environ["SRC"], os.environ["FRAG"], os.environ["OUT"]
FIX1 = os.environ.get("FIX1", "auto")
HERE = os.path.dirname(os.path.abspath(__file__))
LOG = []

# patch-r3.py's helpers, imported rather than copied (its module code only reads the environment)
os.environ.setdefault("L4", SRC)
_spec = importlib.util.spec_from_file_location("patch_r3", os.path.join(HERE, "patch-r3.py"))
r3 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(r3)
title_of, set_notes, copy_slide, move_to, builds, incoming_counts = (r3.title_of, r3.set_notes, r3.copy_slide, r3.move_to, r3.builds,
                                                                     r3.incoming_counts)

FIRST_OLD = "OpenShell architecture and responsibilities"
LAST_OLD = "Limits of OpenShell enforcement"
KEPT = "Real credentials never enter the sandbox."
DIVIDER = "OpenShell and the sandbox research"
OPENER = "OpenShell sandboxes an unmodified agent."
MIDTERM_TITLE = "What the midterm covers"
KEPT_AFTER = "7.10"           # the kept credentials build follows build 10 (provider and credentials)
# A note that names a slide of the old segment names the new build that took over its subject (its last click, as the old numbers
# name a build's complete slide).
TAKES_OVER = {
    "OpenShell architecture and responsibilities": "7.2",       # components
    "OpenShell controls at creation and run time": "7.6",       # the policy's sections pinned to the components
    "OpenShell policy schema": "7.7",                           # the policy, line by line
    "A denied request is recorded for approval.": "7.12",       # a denied request becomes a proposal
    "Sentry watches from outside the host.": "7.9",             # every connection passes the supervisor (its last stage is Sentry)
    "OpenShell governs each sandbox on its own.": "7.14",       # what OpenShell does not cover (Tigera's reading)
    "OpenShell controls by protected asset": "7.14",            # the protected-asset figure
    LAST_OLD: "7.14",
}
NUMREF = re.compile(r"(?i)(?<!removed )\b(slides?)(\s+)(\d+)((?:\s*(?:,|and|to|–|-|&)\s*\d+)*)")
PART = re.compile(r"\s*(,|and|to|–|-|&)\s*(\d+)")


def log(msg):
    LOG.append(msg)
    print(msg)


def clean(t):
    """A title without its reference markers ('...model[1]' -> '...model')."""
    return re.sub(r"(\[\d+\])+$", "", t.split("\n")[0].strip()).strip()


def run_of(prs, title):
    idx = [i for i, s in enumerate(prs.slides) if clean(title_of(s)) == title]
    assert idx and idx == list(range(idx[0], idx[-1] + 1)), (title, idx)
    return idx


# ----------------------------------------------------------------------------------------------- 0. Task 4.3's fix round
def with_fix1(src):
    prs = Presentation(src)
    mid = [i for i, s in enumerate(prs.slides) if clean(title_of(s)) == MIDTERM_TITLE]
    if FIX1 == "0" or len(mid) != 1:
        log(f"fix1: not applied ({'FIX1=0' if FIX1 == '0' else f'the midterm build already has {len(mid)} clicks'})")
        return prs
    tmp = tempfile.mkdtemp(prefix="r6c-fix1-")
    out = os.path.join(tmp, "fix1.pptx")
    env = dict(os.environ, MODE="fix1", SRC=src, OUT=out)
    r = subprocess.run([sys.executable, os.path.join(HERE, "patch-r6.py")], env=env, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[-3000:]
    for line in r.stdout.strip().splitlines():
        log("fix1: " + line)
    return Presentation(out)


# ----------------------------------------------------------------------------------------------- 2. the fragment's slides
def slide_jumps(slide):
    """(rPr ordinal among the slide's a:rPr, child index, hlinkClick element, target part) for every slide-jump link."""
    rprs = list(slide.shapes._spTree.iter(qn("a:rPr")))
    out = []
    for k, rpr in enumerate(rprs):
        for j, h in enumerate(rpr):
            if h.tag == qn("a:hlinkClick") and h.get("action", "").startswith("ppaction://hlinksldjump"):
                out.append((k, j, h, slide.part.rels[h.get(qn("r:id"))].target_part))
    return out


def copy_notes(dst, src):
    """dst's notes become src's, paragraph for paragraph (the fragment's notes are plain runs; the copy keeps their properties)."""
    set_notes(dst, "")
    a, b = src.notes_slide.notes_placeholder._element, dst.notes_slide.notes_placeholder._element
    body = copy.deepcopy(a.find(qn("p:txBody")))
    assert not body.xpath(".//@r:id"), "a relationship in the notes"
    b.replace(b.find(qn("p:txBody")), body)
    assert dst.notes_slide.notes_text_frame.text == src.notes_slide.notes_text_frame.text


def copy_fragment(prs, frag):
    """Every fragment slide but the title slide, appended to prs in order. Returns {fragment position (1-based): new slide}."""
    slides = list(frag.slides)
    assert slides[0].slide_layout.name == "TITLE", slides[0].slide_layout.name
    pos_of = {s.part: k for k, s in enumerate(slides, 1)}
    new, jumps = {}, {}
    for k, fs in enumerate(slides[1:], 2):
        found = slide_jumps(fs)
        for _, _, h, _ in found:                      # copy_slide refuses a slide-jump link; it is put back below
            h.getparent().remove(h)
        c = copy_slide(prs, fs)
        copy_notes(c, fs)
        assert not NUMREF.search(c.notes_slide.notes_text_frame.text), ("a slide number in the fragment's notes", k)
        new[k], jumps[k] = c, found
    relinked = 0
    for k, found in jumps.items():
        c = new[k]
        rprs = list(c.shapes._spTree.iter(qn("a:rPr")))
        for ordinal, j, h, target in found:
            h = copy.deepcopy(h)
            h.set(qn("r:id"), c.part.relate_to(new[pos_of[target]].part, RT.SLIDE))
            rprs[ordinal].insert(j, h)
            relinked += 1
    log(f"fragment: {len(new)} slides copied (fragment slides 2 to {len(slides)}), {relinked} slide-jump links re-pointed at the copies")
    return new


def drop_slides(prs, slides):
    """Remove slides from the deck: the sldId and the presentation's relationship; python-pptx then saves only the parts still reachable.
    No kept slide may link to a dropped one."""
    gone = {s.part for s in slides}
    for s in prs.slides:
        if s.part in gone:
            continue
        for h in s._element.iter(qn("a:hlinkClick")):
            rid = h.get(qn("r:id"))
            assert not (rid and h.get("action", "").startswith("ppaction://hlinksldjump") and s.part.rels[rid].target_part in gone), \
                ("a kept slide links to a dropped one", title_of(s))
    lst = prs.slides._sldIdLst
    for sid in [s.slide_id for s in slides]:
        el = next(e for e in lst if int(e.get("id")) == sid)
        rid = el.get(qn("r:id"))
        lst.remove(el)
        prs.part.drop_rel(rid)
        assert rid not in prs.part.rels, rid


# ----------------------------------------------------------------------------------------------- 3. slide numbers
def renumber(prs, pos_map, old_last, new_last, skip_ids):
    """Link texts follow their targets; every 'slide N' in the notes (but the copied slides') names N's new position."""
    index_of = {s.part: i + 1 for i, s in enumerate(prs.slides)}
    n = len(prs.slides)
    linked, shifted, changes = 0, 0, []
    for s in prs.slides:
        for h in s._element.iter(qn("a:hlinkClick")):
            if h.get("action", "").startswith("ppaction://hlinksldjump"):
                t = h.getparent().getparent().find(qn("a:t"))
                if t is not None and t.text and t.text.strip().isdigit():
                    new = str(index_of[s.part.rels[h.get(qn("r:id"))].target_part])
                    if new != t.text:
                        t.text = new
                        linked += 1

    def one(m):
        word, sp, first, rest = m.groups()
        nums = [(None, int(first))] + [(c, int(d)) for c, d in PART.findall(rest)]
        out = []
        for conn, num in nums:
            assert num in pos_map, ("a slide number this deck does not have", m.group(0))
            new = new_last if conn in ("to", "–", "-") and num == old_last else pos_map[num]
            out.append((conn, new))
        if len({v for _, v in out}) == 1 and len(out) > 1:      # 'slides 129 and 131' now name one build: 'slide N'
            out = out[:1]
            word = word[:-1] if word.lower().endswith("s") else word
        for _, v in out:
            assert 1 <= v <= n, (m.group(0), v)
        text = f"{word}{sp}{out[0][1]}"
        for conn, v in out[1:]:
            text += (", " if conn == "," else f" {conn} ") + str(v)
        return text
    for i, s in enumerate(prs.slides, 1):
        if s.slide_id in skip_ids or not s.has_notes_slide:
            continue
        for t in s.notes_slide.notes_text_frame._txBody.iter(qn("a:t")):
            if t.text and NUMREF.search(t.text):
                new = NUMREF.sub(one, t.text)
                if new != t.text:
                    for a, b in zip(NUMREF.finditer(t.text), NUMREF.finditer(new)):
                        if a.group(0) != b.group(0):
                            changes.append((i, a.group(0), b.group(0)))
                    t.text = new
                    shifted += 1
    log(f"renumbered: {linked} link texts, {shifted} note runs")
    seen = {}
    for i, a, b in changes:
        seen.setdefault((a, b), []).append(i)
    for (a, b), where in seen.items():
        log(f"  notes {where[0]}-{where[-1]}: {a!r} -> {b!r}")


# ----------------------------------------------------------------------------------------------- 4. back links
def drop_orphan_back_links(prs, before):
    """lectures/CLAUDE.md: the '‹ back' link sits on every click of a build that a slide number points to. A build whose only incoming
    links came from the old segment loses it."""
    after = incoming_counts(prs)
    slides = list(prs.slides)
    removed = []
    for b in builds(prs):
        key = slides[b[0]].slide_id
        if before.get(key, 0) > 0 and after.get(key, 0) == 0:
            for i in b:
                for sh in list(slides[i].shapes):
                    if sh.name == "Back link":
                        sh._element.getparent().remove(sh._element)
                        removed.append(i + 1)
    log(f"back links: {len(removed)} removed from builds that lost their last incoming link: {removed}")


# ----------------------------------------------------------------------------------------------- part names
def unique_partnames(prs, new_parts):
    """python-pptx names an added slide slide<count + 1>.xml, which can be the name of a slide the deck already holds (lectures/CLAUDE.md:
    'deleting and adding slides in one run can reuse a name a surviving slide still holds'). A new slide whose name another part holds
    takes the next free number."""
    from pptx.opc.packuri import PackURI
    parts = list(prs.part.package.iter_parts())
    names = [str(p.partname) for p in parts]
    taken = set(names)
    nums = [int(m.group(1)) for n in names for m in [re.fullmatch(r"/ppt/slides/slide(\d+)\.xml", n)] if m]
    nxt = max(nums) + 1
    renamed = 0
    for p in new_parts:
        if names.count(str(p.partname)) > 1:
            while f"/ppt/slides/slide{nxt}.xml" in taken:
                nxt += 1
            names[names.index(str(p.partname))] = None
            p.partname = PackURI(f"/ppt/slides/slide{nxt}.xml")
            taken.add(str(p.partname))
            names.append(str(p.partname))
            renamed += 1
    final = [str(p.partname) for p in prs.part.package.iter_parts()]
    assert len(final) == len(set(final)), sorted({n for n in final if final.count(n) > 1})
    log(f"part names: {renamed} new slides renamed to free names, no duplicate part names")


# ----------------------------------------------------------------------------------------------- main
def main():
    prs = with_fix1(SRC)
    frag = Presentation(FRAG)
    fmap = json.load(open(FRAG + ".map.json"))["positions"]
    n0 = len(prs.slides)
    before_links = incoming_counts(prs)
    before = [s.part for s in prs.slides]
    titles = [clean(title_of(s)) for s in prs.slides]

    # 1. the old segment
    first = titles.index(FIRST_OLD)
    last = len(titles) - 1 - titles[::-1].index(LAST_OLD)
    old = list(prs.slides)[first:last + 1]
    kept_idx = run_of(prs, KEPT)
    assert first < kept_idx[0] and kept_idx[-1] < last, (first, kept_idx, last)
    kept = [prs.slides[i] for i in kept_idx]
    assert titles[first - 1] == OPENER and DIVIDER in titles[:first], (titles[first - 1],)
    old_titles = sorted({titles[i] for i in range(first, last + 1)} - {KEPT})
    assert set(old_titles) == set(TAKES_OVER), ("the old segment's builds are not the expected ones", old_titles)
    log(f"old segment: slides {first + 1}-{last + 1} ({len(old)} slides, {len(builds_in(titles, first, last))} builds); "
        f"kept: {KEPT!r} at {kept_idx[0] + 1}-{kept_idx[-1] + 1}")

    # 2. the fragment, the kept build after build 10, the new run in the old segment's place, the old slides dropped
    new = copy_fragment(prs, frag)
    order = sorted(new)
    after10 = fmap[KEPT_AFTER][-1]
    run = [new[k] for k in order if k <= after10] + kept + [new[k] for k in order if k > after10]
    move_to(prs, run, first)
    drop_slides(prs, [s for s in old if s not in kept])
    slides = list(prs.slides)
    assert slides[first:first + len(run)] == run
    new_first, new_last = first + 1, first + len(run)

    # 3. slide numbers: old position -> new position
    position = {s.part: i + 1 for i, s in enumerate(slides)}
    last_of = {key: position[new[fmap[key][-1]].part] for key in fmap}
    pos_map = {}
    for k, part in enumerate(before, 1):
        pos_map[k] = position[part] if part in position else last_of[TAKES_OVER[titles[k - 1]]]
    renumber(prs, pos_map, last + 1, new_last, {s.slide_id for s in new.values()})

    # 4. back links
    drop_orphan_back_links(prs, before_links)

    # 5. save
    unique_partnames(prs, [s.part for s in new.values()])
    want = n0 - len(old) + len(new) + len(kept)
    assert len(prs.slides) == want, (len(prs.slides), want)
    log(f"new segment: slides {new_first}-{new_last} ({len(run)} slides: {len(new)} from the fragment, {len(kept)} kept)")
    for key in fmap:
        a, b = position[new[fmap[key][0]].part], position[new[fmap[key][-1]].part]
        log(f"  build {key}: {a}-{b} {clean(title_of(slides[a - 1]))!r}")
        if key == KEPT_AFTER:
            log(f"  build 7.11 (kept): {position[kept[0].part]}-{position[kept[-1].part]} {KEPT!r}")
    prs.save(OUT)
    log(f"saved {OUT}: {n0} -> {len(prs.slides)} slides ({n0} - {len(old)} old + {len(new)} fragment + {len(kept)} kept)")
    with open(OUT + ".log", "w") as f:
        f.write("\n".join(LOG) + "\n")


def builds_in(titles, first, last):
    out, prev = [], None
    for i in range(first, last + 1):
        if titles[i] != prev:
            out.append(i)
        prev = titles[i]
    return out


if __name__ == "__main__":
    main()

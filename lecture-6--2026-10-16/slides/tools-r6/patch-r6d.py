"""Remove the instructor's name from lecture6.pptx: the Cyberbird bullet of the gateway build (slides 116 to 121), the
notes of the components build (slides 101 to 108), and the "in action" stage of the components and gateway-authority
figures, swapped for the re-rendered stages in slides/figures/r6c. Slide 120 already read "our client key" (the
instructor's own edit), so every replacement uses that wording.

    cd lecture-6--2026-10-16 && IN=lecture6.pptx OUT=<scratch>/lecture6.r6d.pptx OLD=<folder with the old two PNGs> \
        uv run --no-project --python 3.12 --with python-pptx python slides/tools-r6/patch-r6d.py
"""
import hashlib, os
from pptx import Presentation

IN, OUT, OLD = os.environ["IN"], os.environ["OUT"], os.environ["OLD"]
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "figures", "r6c")
TEXT = {"Rahul's client key": "our client key", "Rahul's terminal": "our terminal"}
STAGES = ("components-7.png", "gateway-authority-5.png")

sha = lambda b: hashlib.sha1(b).hexdigest()
swap = {sha(open(os.path.join(OLD, n), "rb").read()): open(os.path.join(FIG, n), "rb").read() for n in STAGES}
for n in STAGES:
    assert open(os.path.join(OLD, n), "rb").read() != open(os.path.join(FIG, n), "rb").read(), f"{n} was not re-rendered"

p = Presentation(IN)
runs = pictures = 0
for i, s in enumerate(p.slides, 1):
    frames = [sh.text_frame for sh in s.shapes if sh.has_text_frame]
    if s.has_notes_slide:
        frames.append(s.notes_slide.notes_text_frame)
    for tf in frames:
        for para in tf.paragraphs:
            for r in para.runs:
                for old, new in TEXT.items():
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        runs += 1
    for sh in s.shapes:
        if sh.shape_type == 13:  # picture
            img = sh._parent.part.related_part(sh._element.blip_rId)
            h = sha(img.blob)
            if h in swap:
                img._blob = swap[h]
                pictures += 1

left = [(i, "Rahul") for i, s in enumerate(p.slides, 1)
        for tf in [sh.text_frame for sh in s.shapes if sh.has_text_frame]
                  + ([s.notes_slide.notes_text_frame] if s.has_notes_slide else [])
        if "Rahul" in tf.text]
print(f"runs replaced {runs}, picture parts swapped {pictures}, mentions left {left}")
assert runs == 13 and not left, "unexpected counts: the deck changed since the patch was written"
assert pictures >= 1, "no figure stage matched the old images"
p.save(OUT)
print("wrote", OUT)

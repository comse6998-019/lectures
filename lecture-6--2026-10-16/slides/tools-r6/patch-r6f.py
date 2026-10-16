"""Merge the interleaved components and lifecycle builds of lecture6.pptx (slides 101 to 115 after patch-r6e.py) into one
build of 13 clicks over one figure, slides/figures/r6c/components-lifecycle-1.png to -12.png: the components drawing with the
sandbox's phases in a band above it, each phase arriving on the click after the component that performs it, then the closing
box. Every slide of the build is rebuilt from the components build's closing slide (its title, body, closing box and
reference line); the lifecycle's points that the figure has no room for (Error, Deleting and Stopped, and Ready saying nothing
about the agent) move to the notes. 15 slides become 13, so every later slide moves up by two; no link points into the range.

    cd lecture-6--2026-10-16 && IN=lecture6.pptx OUT=<scratch>/lecture6.r6f.pptx \
        uv run --no-project --python 3.12 --with python-pptx,pillow python slides/tools-r6/patch-r6f.py
"""
import copy, os
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

IN, OUT = os.environ["IN"], os.environ["OUT"]
HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "..", "figures", "r6c", "components-lifecycle-{}.png")
COMP, LIFE = "Each OpenShell component has its own job.", "A sandbox is a record before a container."
INTERLEAVED = "CCLCLCCLLLCCCLL"            # slides 101 to 115 after patch-r6e.py
COL_L, COL_W, BODY_TOP, GAP = 320040, 6016752, 658368, 186667
BOX_TOP, UNDER = 4370832, 251381          # the one-line closing box, 251381 under the figure
BODY_H = 870000                           # three one-line bullets at 12 pt with 3 pt before and after

BULLETS = [("The gateway", " records each sandbox before the driver provisions it."),
           ("The sandbox runtime", " fences the agent's processes from inside the workload."),
           ("The trusted supervisor", " connects back, then decides every connection.")]
# (figure stage, bullets shown, closing box) for each click
CLICKS = [(s, 1 if s <= 5 else 2 if s == 6 else 3, False) for s in range(1, 13)] + [(12, 3, True)]

NOTES = """We name OpenShell's parts and follow one sandbox through them. The figure is the guide's Figure 1 with the sandbox's phases in a band above it, and each phase arrives on the click after the component that performs it. The client (the openshell CLI, the TUI or an SDK) sends management requests to OpenShell's gateway, the control plane that authenticates every caller, stores every sandbox record and hands configuration to drivers and supervisors. The gateway checks a create request (identity, workspace, names, template, policy) and persists the record before any container exists, with a UUID, an attachment epoch that stops a re-attached provider from reusing old evidence, and a runtime generation that stops a new run from inheriting old runtime credentials. The compute driver turns the record into platform resources on Docker, Podman, Kubernetes or a microVM and decides placement, which is on which host or node the workload and its supervisor run. Docker and the VM may answer early and keep preparing the image, Kubernetes reconciles, and a watch stream reports progress. The untrusted workload is the agent with its tools and child processes, code we assume may be faulty or hostile, together with the sandbox runtime, the small OpenShell program inside it that launches processes, installs the kernel fences and pauses network calls for mediation. The trusted supervisor runs outside that boundary, holds the policy and the real credentials, decides each outbound connection, and talks to the runtime over the authenticated runtime channel, a private connection whose ends authenticate each other. The supervisor connects out to the gateway and says hello, and the gateway then publishes the sandbox as Ready, usable through its API.

Ready says the supervisor talked back. It says nothing about whether the agent works, and the phases Error (a backend failure the driver reports), Deleting and Stopped still outrank it. Stop keeps the record and stops the compute, start runs it again, and delete removes it, with what is retained depending on the backend: Docker keeps the stopped container, Kubernetes its Sandbox resource, the VM its overlay. The gateway records each phase, so an interrupted create can be reconciled.

On Docker the workload and the supervisor are separate containers joined by a private Unix socket, the workload's own network is removed, and the supervisor uses the host's network (guide, 6.3). In the recorded run the client was the instructor's terminal, the driver Docker on Colima, and the policy admitted one destination, the model server. Captured for Cyberbird on October 15, 2026 (OpenShell 0.1.4, Colima): create printed "Created sandbox: cyberbird", "[0.0s] Requesting compute...", "[0.0s] Sandbox allocated" and "[0.1s] Starting sandbox Container created"; sandbox get showed "Ready: True (DependenciesReady) - Supervisor session connected", Id 86d1c2e9-99a7-4c49-9101-118c7f4e3230 and runtime-generation d1326079-66a5-44ab-9db3-51e13ac8dc46; sandbox list at 22:42:24 EDT printed "cyberbird  2026-10-16 02:42:20  Ready" (CREATED is UTC, so 22:42:20 EDT).

The figure, stage by stage.
1. The client: the CLI, TUI or an SDK sends management requests. The agent and its tools in the untrusted workload are the frame, and the sandbox's phases wait above as ghosts.
2. The gateway: identity and records, and the configuration it hands on.
3. The record persisted, by the gateway: its checks pass and it stores the record, in the Provisioning phase, before any container exists.
4. The compute driver provisions the workload and places the supervisor. The note defines placement.
5. Provisioning, by the driver: it creates the resources. Docker and the VM may return early.
6. The sandbox runtime, inside the workload: process control and mediation.
7. The trusted supervisor, outside: policy and credentials, the authenticated runtime channel from the runtime, and the control session it opens to the gateway.
8. The supervisor's hello: it connects out to the gateway and opens the control stream that later carries exec's relay.
9. Ready: the gateway publishes the sandbox as usable through its API. Ready says the supervisor talked back, and nothing about the agent.
10. The approved external service: what remains reachable, through the supervisor.
11. In action: our terminal, Docker on Colima, cyberbird and bandit in the workload, and a supervisor that admits only the model server.
12. In action, the phases: cyberbird's record and UUID, Docker on Colima with the image cyberbird:1, the captured condition "Supervisor session connected", and sandbox list showing Ready."""

p = Presentation(IN)
slides = list(p.slides)
title = lambda s: s.shapes.title.text_frame.text if s.shapes.title is not None else ""
assert len(slides) == 222
got = "".join("C" if title(slides[n - 1]).startswith(COMP) else "L" if title(slides[n - 1]).startswith(LIFE) else "?" for n in range(101, 116))
assert got == INTERLEAVED, f"slides 101 to 115 are {got}, expected the interleaved {INTERLEAVED}"
assert not title(slides[99]).startswith((COMP, LIFE)) and not title(slides[115]).startswith((COMP, LIFE))
template = slides[112]                      # the components closing slide: title, body, closing box, reference line
assert any(sh.name == "Rectangle 4" for sh in template.shapes)
named = lambda s, n: [sh for sh in s.shapes if sh.name == n][0]
vis_par = [copy.deepcopy(named(template, "Content Placeholder 2").text_frame.paragraphs[k]._p) for k in range(3)]
hid_par = copy.deepcopy(named(slides[100], "Content Placeholder 2").text_frame.paragraphs[1]._p)   # a hidden bullet
assert hid_par.find(".//" + qn("a:buNone")) is not None
keep_shapes = [copy.deepcopy(sh._element) for sh in template.shapes if sh.shape_type != 13]
pic_name = [sh.name for sh in template.shapes if sh.shape_type == 13][0]
cite = named(template, "Citations").text_frame.text
assert cite.endswith("chapters 2, 6 and 7."), cite

def bullet(k, shown):
    par = copy.deepcopy(vis_par[k] if shown else hid_par)
    runs = par.findall(qn("a:r"))
    assert len(runs) == 2
    runs[0].find(qn("a:t")).text, runs[1].find(qn("a:t")).text = BULLETS[k]
    return par

w_px, h_px = Image.open(FIG.format(1)).size
assert all(Image.open(FIG.format(k)).size == (w_px, h_px) for k in range(1, 13)), "stages differ in size"
fig_top = BODY_TOP + BODY_H + GAP
fig_h = BOX_TOP - UNDER - fig_top
fig_w = round(fig_h * w_px / h_px)
assert 0.9 * COL_W <= fig_w <= COL_W, f"figure at {fig_w / COL_W:.0%} of the column"
fig_l = COL_L + (COL_W - fig_w) // 2

targets, extra = slides[100:113], slides[113:115]
for (stage, nb, box), s in zip(CLICKS, targets):
    tree = s.shapes._spTree
    for el in list(tree)[2:]:
        tree.remove(el)
    for rel in [r for r in s.part.rels.values() if r.reltype.endswith("/image")]:
        s.part.drop_rel(rel.rId)
    for el in keep_shapes:
        tree.append(copy.deepcopy(el))
    if not box:
        tree.remove(named(s, "Rectangle 4")._element)
    body = named(s, "Content Placeholder 2")
    body.height = Emu(BODY_H)
    txb = body.text_frame._txBody
    for par in txb.findall(qn("a:p")):
        txb.remove(par)
    for k in range(3):
        txb.append(bullet(k, k < nb))
    named(s, "Citations").text_frame.paragraphs[0].runs[0].text = cite.replace("chapters 2, 6 and 7.", "chapters 2, 4, 6 and 7.")
    pic = s.shapes.add_picture(FIG.format(stage), Emu(fig_l), Emu(fig_top), Emu(fig_w), Emu(fig_h))
    pic.name = pic_name
    s.notes_slide.notes_text_frame.text = NOTES

sld_ids = p.slides._sldIdLst
for s in extra:
    for sid in list(sld_ids):
        if p.part.related_part(sid.get(qn("r:id"))) is s.part:
            p.part.drop_rel(sid.get(qn("r:id")))
            sld_ids.remove(sid)
p.save(OUT)

q = Presentation(OUT)
assert len(q.slides) == 220
run = [title(q.slides[n - 1]) for n in range(100, 116)]
assert all(t.startswith(COMP) for t in run[1:14]) and not run[0].startswith(COMP) and not run[14].startswith((COMP, LIFE))
print(f"220 slides; the build is slides 101 to 113; figure {fig_w / COL_W:.1%} of the column, {fig_l},{fig_top} {fig_w}x{fig_h}")

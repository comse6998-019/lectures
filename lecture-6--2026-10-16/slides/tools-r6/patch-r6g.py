"""Bring the command x-ray forward in lecture6.pptx. The build "The commands HW2 uses." (slides 180 to 185 after patch-r6f.py)
splits at its figure's stage 3: sandbox create and sandbox exec (stages 1 to 3) move to directly after the components build
(slides 101 to 113), so each token of a student's command meets the components and phases it sets in motion; logs and policy
get (stages 4 and 5) stay in place, after the policy, because they show the supervisor's decisions. The figure is
slides/figures/r6c/command-hw2-1.png to -5.png, unchanged, one command per stage; each half keeps only its own two bullets, so the figure moves up
under them and widens to the column, and each half ends on its own closing box. 6 slides become 7.

    cd lecture-6--2026-10-16 && IN=lecture6.pptx OUT=<scratch>/lecture6.r6g.pptx \
        uv run --no-project --python 3.12 --with python-pptx,pillow python slides/tools-r6/patch-r6g.py
"""
import copy, os
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

IN, OUT = os.environ["IN"], os.environ["OUT"]
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "figures", "r6c", "command-hw2-1.png")
OLD_TITLE, COMP = "The commands HW2 uses.", "Each OpenShell component has its own job."
FIRST_TITLE = "Create and exec are requests to the gateway."
SECOND_TITLE = "Logs and policy get show OpenShell's decisions."
FIRST_CLOSE = "Create asks the gateway for a sandbox, and exec runs a command inside it."
SECOND_CLOSE = "Logs show each decision, and policy get shows the policy behind it."
OLD_CLOSE = "Flags before the double dash are the operator's."
COL_L, COL_W, BODY_TOP, GAP, UNDER, BOX_H = 320040, 6016752, 658368, 186667, 251381, 429768
BODY_H = 600000                                   # two one-line bullets at 12 pt

FIRST_NOTES = """We read the two commands that start a sandbox and run a workload in it, token by token, against the components and phases of the previous build. sandbox create asks the gateway for a sandbox: the gateway persists the record, the driver provisions the container, and the supervisor's hello makes it Ready. sandbox exec asks the gateway to run one more command in that sandbox: the gateway allocates a relay channel and asks the supervisor to connect back, and output and exit status travel back as API events (guide, 4.5). Everything after the double dash is the workload's command, and everything before it is the operator's, which OpenShell reads. sandbox create and sandbox exec appear in the brief's README section, "The sandbox image and the launch commands", and exec again in "How the guardrails are tested": the staff driver starts each team's MCP server with exec --no-login-shell --no-tty, because a login shell's start-up files can print to standard output and corrupt the JSON-RPC stream, whose first byte must be {. logs and policy get, the other two commands HW2 uses, come after the policy.

The commands ran on October 15, 2026 with OpenShell 0.1.4 (cyberbird/create.txt, run.txt): create idled the sandbox on sleep infinity and exec started Cyberbird. exec also passed --env OPENAI_BASE_URL and --env OPENAI_API_KEY, because 0.1.4 hands exec none of the image's ENV values except PATH. openshell sandbox list then printed "cyberbird  2026-10-16 02:42:20  Ready".

The figure, stage by stage.
1. The shape of sandbox create: a name, an image, a policy, an optional provider, and after the double dash the workload's command.
2. sandbox create for Cyberbird, and the lines it printed: the gateway's record, then the container.
3. sandbox exec: the sandbox, no login shell, no terminal, and Cyberbird's command, which runs through the supervisor's relay, with the two lines the run ended on."""

SECOND_NOTES = """We return to the commands HW2 uses: logs and policy get show what OpenShell decided under a policy. logs --source sandbox prints the sandbox's own events, its connection decisions, policy loads and process events, and never the main process's ordinary output (guide, 13.4). policy get --full prints the effective policy, the authored document plus the rules that attached providers add (guide, 10.4). policy get --full appears in the brief's "The guardrails", where A4 asks for a reason for each grant, and the driver reads it to confirm the provider is the only network grant; logs is ours, the evidence of a denial. Each stage of the figure shows one command: create and exec were taken apart in the build after the components.

The commands ran on October 15, 2026 with OpenShell 0.1.4 (cyberbird/logs.txt, policy-full.txt).

The figure, stage by stage.
1. logs: the sandbox's own events. python3.12 is denied pypi.org:443, curl is denied the model server, and python3.12 is allowed to it: a rule admits a binary together with a destination.
2. policy get --full: the effective policy. With no provider attached it holds no provider rules."""

p = Presentation(IN)
slides = list(p.slides)
title = lambda s: s.shapes.title.text_frame.text if s.shapes.title is not None else ""
assert len(slides) == 220
assert all(title(slides[n - 1]) == OLD_TITLE for n in range(180, 186)), "slides 180 to 185 are not the command build"
assert title(slides[178]) != OLD_TITLE and title(slides[185]) != OLD_TITLE
assert all(title(slides[n - 1]).startswith(COMP) for n in range(101, 114)) and not title(slides[113]).startswith(COMP)
named = lambda s, n: [sh for sh in s.shapes if sh.name == n][0]
close_src = named(slides[184], "Rectangle 4")
assert close_src.text_frame.text.startswith(OLD_CLOSE)

def clone(src):
    new = p.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree)[2:]:
        tree.remove(el)
    rid = {}
    for old_id, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        rid[old_id] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external \
            else new.part.relate_to(rel.target_part, rel.reltype)
    for el in list(src.shapes._spTree)[2:]:
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % R) and v in rid:
                el.set(k, rid[v])
    return new

w_px, h_px = Image.open(FIG).size
fig_top = BODY_TOP + BODY_H + GAP
fig_h = round(COL_W * h_px / w_px)
box_top = fig_top + fig_h + UNDER
assert box_top + BOX_H <= 4800600, "the closing box would cross the floor"

def shape_up(s, ttl, drop, close, notes):
    s.shapes.title.text_frame.paragraphs[0].runs[0].text = ttl
    body = named(s, "Content Placeholder 2")
    txb = body.text_frame._txBody
    pars = txb.findall(qn("a:p"))
    assert len(pars) == 4
    for k in drop:
        txb.remove(pars[k])
    body.height = Emu(BODY_H)
    pic = [sh for sh in s.shapes if sh.shape_type == 13][0]
    pic.left, pic.top, pic.width, pic.height = Emu(COL_L), Emu(fig_top), Emu(COL_W), Emu(fig_h)
    boxes = [sh for sh in s.shapes if sh.name == "Rectangle 4"]
    if close:
        if not boxes:
            s.shapes._spTree.append(copy.deepcopy(close_src._element))
            boxes = [sh for sh in s.shapes if sh.name == "Rectangle 4"]
        box = boxes[0]
        par = box.text_frame.paragraphs[0]
        par.runs[0].text = close
        for r in par.runs[1:]:
            par._p.remove(r._r)
        box.top, box.height = Emu(box_top), Emu(BOX_H)
    else:
        assert not boxes
    s.notes_slide.notes_text_frame.text = notes

first_close = clone(slides[181])                 # stage 3 with both bullets, plus the closing box
first = [slides[179], slides[180], slides[181], first_close]
second = [slides[182], slides[183], slides[184]]
for k, s in enumerate(first):
    shape_up(s, FIRST_TITLE, (3, 2), FIRST_CLOSE if k == 3 else None, FIRST_NOTES)
for k, s in enumerate(second):
    shape_up(s, SECOND_TITLE, (1, 0), SECOND_CLOSE if k == 2 else None, SECOND_NOTES)

lst = p.slides._sldIdLst
ids = list(lst)
by_part = {p.part.related_part(sid.get(qn("r:id"))): sid for sid in ids}
anchor = by_part[slides[112].part]               # the components build's closing slide
for s in reversed(first):
    el = by_part[s.part]
    lst.remove(el)
    anchor.addnext(el)
p.save(OUT)

q = Presentation(OUT)
t = [title(s) for s in q.slides]
assert len(t) == 221
assert t[113:117] == [FIRST_TITLE] * 4 and t[112].startswith(COMP) and not t[117].startswith(FIRST_TITLE)
i = t.index(SECOND_TITLE)
assert t[i:i + 3] == [SECOND_TITLE] * 3 and OLD_TITLE not in t
print(f"221 slides; create and exec at 114 to 117, logs and policy get at {i + 1} to {i + 3}; figure top {fig_top}, "
      f"height {fig_h}, closing box at {box_top}")

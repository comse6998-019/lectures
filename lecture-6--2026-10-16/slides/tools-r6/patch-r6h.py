"""Build Cyberbird's policy block by block in lecture6.pptx. The build "Cyberbird's policy, line by line" (slides 137 to 140 after
patch-r6g.py) showed the whole file in one code window. It becomes "Cyberbird's policy, block by block": one click per block of
the policy over the staged figure slides/figures/r6c/policy-blocks-1.png to -8.png, which sets each block beside what it grants
(the files, as Landlock enforces them; the connections, as the supervisor decides them), then the closing box. The first
bullet ("As a document, ...") moves to the notes; the other two arrive with the file lists and the binaries. 4 slides become 9.

    cd lecture-6--2026-10-16 && IN=lecture6.pptx OUT=<scratch>/lecture6.r6h.pptx \
        uv run --no-project --python 3.12 --with python-pptx,pillow python slides/tools-r6/patch-r6h.py
"""
import copy, os
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu

IN, OUT = os.environ["IN"], os.environ["OUT"]
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "figures", "r6c", "policy-blocks-{}.png")
OLD, NEW = "Cyberbird's policy, line by line", "Cyberbird's policy, block by block"
CLOSE = "Each grant covers everything beneath its path, the answer key included."
WINDOW = {"CodeWindow", "WinBar", "WinDot 0", "WinDot 1", "WinDot 2", "TextBox 8", "TextBox 7"}
COL_L, COL_W, BODY_TOP, GAP, UNDER, BOX_H = 320040, 6016752, 658368, 186667, 251381, 429768
BODY_H = 600000
# (figure stage, bullets shown, closing box) per click: the read list arrives with the file lists, the network rule with the binaries
CLICKS = [(1, 0, False)] + [(s, 1, False) for s in (2, 3, 4, 5)] + [(s, 2, False) for s in (6, 7, 8)] + [(8, 2, True)]

NOTES = """We read Cyberbird's policy one block at a time and watch what each block grants. The text is the policy of the run we recorded on October 15, 2026, instructor-materials/hw2/feasibility/cyberbird-policy.yaml, with its two lists wrapped to fit and its three comment lines left out; the comments say why the host is host.lima.internal: OpenShell 0.1.4 on Colima refuses host.docker.internal at policy DNS, and an IP literal is refused for protocol tcp. On the right, the files the workload can touch and the connections it can open, with the outcomes captured inside the sandbox that evening (probes.txt and logs.txt).

As a document, a policy is parsed, validated, versioned, diffed, loaded and proved: a bounded parser rejects unknown fields, duplicate keys and oversized input, the gateway validates and versions every revision, a reviewer diffs it, the supervisor loads it and the prover reads it. The read list grants /app, so inside the sandbox the sibling case BenchmarkTest00284.py and the answer key, /app/fixtures/owasp-benchmark-python/ground-truth/expectedresults-0.1.csv, were both readable; only Cyberbird's own workspace code kept them out of the agent's copy, and the HW2 policy lists the target's code directories instead. A grant covers everything beneath its path, so /app/runs is readable from the first block and writable once read_write names it. protocol: tcp admits the connection with no request inspection, so the rule has no enforcement line; with --model anthropic:claude-opus-5 the rule would name api.anthropic.com port 443 with protocol: rest and enforcement: enforce, and the key would come through a provider attachment. The effective policy the supervisor loaded, openshell policy get cyberbird --full, is version 2: OpenShell added /var/log to read_only and the name model_server to the rule.

The figure, stage by stage.
1. The frame: the policy's version line, and the files and connections with nothing granted yet.
2. filesystem_policy and its read_only list: everything under /bin, /usr, /lib, /etc, /proc, /dev/urandom and /app becomes readable, /app/runs included.
3. read_write: /sandbox, /tmp, /dev/null and /app/runs become writable. Inside the sandbox a write into the fixture failed and writes into /app/runs and /tmp succeeded.
4. landlock: the kernel holds both lists, and a path neither list names is refused: /home/agent and /var/lib/dpkg/status gave Permission denied. hard_requirement means a sandbox whose kernel cannot enforce the lists does not start; best_effort would continue with a finding.
5. network_policies: one endpoint, the model server at host.lima.internal port 8080, TCP, behind the trusted supervisor.
6. binaries: only /usr/local/bin/python3.12 may use it. The log read NET:OPEN ALLOWED /usr/local/bin/python3.12 -> host.lima.internal:8080 [policy:model_server].
7. What the rule leaves out: curl to the same server was denied, "binary '/usr/bin/curl' not allowed in policy 'model_server'", and python3.12 to pypi.org:443 was denied, "endpoint pypi.org:443 is not allowed by any policy". A rule admits a binary together with a destination.
8. The answer key sits under /app, so the read list grants it: the probe printed its first line."""

p = Presentation(IN)
slides = list(p.slides)
title = lambda s: s.shapes.title.text_frame.text if s.shapes.title is not None else ""
assert len(slides) == 221
assert all(title(slides[n - 1]).startswith(OLD) for n in range(137, 141)), [title(slides[n - 1]) for n in range(137, 141)]
assert not title(slides[135]).startswith(OLD) and not title(slides[140]).startswith(OLD)
named = lambda s, n: [sh for sh in s.shapes if sh.name == n][0]
body_of = lambda s: named(s, "Content Placeholder 2").text_frame.paragraphs
hidden = [copy.deepcopy(body_of(slides[136])[k]._p) for k in (1, 2)]
shown = [copy.deepcopy(body_of(slides[138])[k]._p) for k in (1, 2)]
assert all(h.find(".//" + qn("a:buNone")) is not None for h in hidden) and all(v.find(".//" + qn("a:buNone")) is None for v in shown)
assert named(slides[139], "Rectangle 4").text_frame.text == CLOSE

def clone(src):
    new = p.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree)[2:]:
        tree.remove(el)
    rid = {}
    for old_id, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        rid[old_id] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external \
            else new.part.relate_to(rel.target_part, rel.reltype)
    for el in list(src.shapes._spTree)[2:]:
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % R) and v in rid:
                el.set(k, rid[v])
    return new

w_px, h_px = Image.open(FIG.format(1)).size
assert all(Image.open(FIG.format(k)).size == (w_px, h_px) for k in range(1, 9))
fig_top = BODY_TOP + BODY_H + GAP
fig_h = round(COL_W * h_px / w_px)
box_top = fig_top + fig_h + UNDER
assert box_top + BOX_H <= 4800600, f"closing box ends at {box_top + BOX_H}"

template = slides[139]                               # the old build's closing slide
targets = slides[136:140] + [clone(template) for _ in range(5)]
for (stage, nb, box), s in zip(CLICKS, targets):
    tree = s.shapes._spTree
    for sh in list(s.shapes):
        if sh.name in WINDOW or sh.shape_type == 13:
            tree.remove(sh._element)
    has_box = any(sh.name == "Rectangle 4" for sh in s.shapes)
    if box and not has_box:
        tree.append(copy.deepcopy(named(template, "Rectangle 4")._element))
    if not box and has_box:
        tree.remove(named(s, "Rectangle 4")._element)
    if box:
        b = named(s, "Rectangle 4")
        b.top, b.height = Emu(box_top), Emu(BOX_H)
    s.shapes.title.text_frame.paragraphs[0].runs[0].text = NEW
    body = named(s, "Content Placeholder 2")
    body.height = Emu(BODY_H)
    txb = body.text_frame._txBody
    for par in txb.findall(qn("a:p")):
        txb.remove(par)
    for k in range(2):
        txb.append(copy.deepcopy(shown[k] if k < nb else hidden[k]))
    s.shapes.add_picture(FIG.format(stage), Emu(COL_L), Emu(fig_top), Emu(COL_W), Emu(fig_h)).name = "Picture policy-blocks"
    s.notes_slide.notes_text_frame.text = NOTES

lst = p.slides._sldIdLst
by_part = {p.part.related_part(sid.get(qn("r:id"))): sid for sid in lst}
anchor = by_part[slides[139].part]
for s in reversed(targets[4:]):
    el = by_part[s.part]
    lst.remove(el)
    anchor.addnext(el)
p.save(OUT)

q = Presentation(OUT)
t = [title(s) for s in q.slides]
assert len(t) == 226 and all(x.startswith(NEW) for x in t[136:145]) and not t[145].startswith(NEW) and not t[135].startswith(NEW)
print(f"226 slides; the build is slides 137 to 145; figure top {fig_top}, height {fig_h}, closing box at {box_top}")

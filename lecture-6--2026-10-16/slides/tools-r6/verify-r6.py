"""SRC=<Rahul's save> L5=<416-slide lecture5.pptx> uv run --no-project --python 3.12 --with python-pptx,lxml python verify-r6.py
Maps every slide of SRC to one slide of L5 by (title, notes text) in order; fails on an unmapped or doubly mapped slide."""
import os, sys
from pptx import Presentation
def key(s):
    t = s.shapes.title.text_frame.text.strip() if s.shapes.title is not None else ""
    if not t:
        t = next((sh.text_frame.text.strip().split("\n")[0] for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "")
    n = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
    return (t, n)
src, l5 = Presentation(os.environ["SRC"]), Presentation(os.environ["L5"])
old = [key(s) for s in l5.slides]
mapping, cursor = {}, 0
for i, s in enumerate(src.slides, 1):
    k = key(s)
    if i <= 3 and k[0] in ("Designing Production Agentic Systems", "Lecture 5 Recap"):
        continue                                   # his new front matter
    j = next((j for j in range(cursor, len(old)) if old[j] == k), None)
    assert j is not None, f"slide {i} {k[0]!r} has no match in the 416-slide deck after old {cursor + 1}"
    mapping[i] = j + 1; cursor = j + 1
assert len(set(mapping.values())) == len(mapping), "two slides map to one original"
print(len(mapping), "slides mapped; first", min(mapping.items()), "last", max(mapping.items()))
import json; json.dump(mapping, open(os.environ.get("MAP", "/tmp/map.json"), "w"))

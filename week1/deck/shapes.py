"""Drawn figure primitives. Positions are inches; nothing here knows content."""

from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from deck import theme


def _style_runs(frame, pt, name, color, bold=False):
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(pt)
            run.font.name = name
            run.font.color.rgb = color
            run.font.bold = bold


@dataclass
class Table:
    rows: list
    left: float
    top: float
    width: float
    height: float
    header: bool = True
    pt: float = 10.0

    @property
    def text_for_check(self):
        return "\n".join("\t".join(str(c) for c in row) for row in self.rows)

    def draw(self, slide):
        n_rows, n_cols = len(self.rows), len(self.rows[0])
        graphic = slide.shapes.add_table(
            n_rows,
            n_cols,
            Inches(self.left),
            Inches(self.top),
            Inches(self.width),
            Inches(self.height),
        )
        table = graphic.table
        for r, row in enumerate(self.rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                bold = self.header and r == 0
                _style_runs(
                    cell.text_frame,
                    self.pt,
                    theme.BODY_FONT,
                    theme.INK,
                    bold=bold,
                )
        return graphic


@dataclass
class Boxes:
    labels: list
    left: float
    top: float
    width: float
    height: float
    arrows: bool = True
    fill: object = None
    pt: float = 11.0

    @property
    def text_for_check(self):
        return "\n".join(str(label) for label in self.labels)

    def draw(self, slide):
        count = len(self.labels)
        gap = 0.18 if self.arrows and count > 1 else 0.08
        box_w = (self.width - gap * (count - 1)) / count
        drawn = []
        for i, label in enumerate(self.labels):
            x = self.left + i * (box_w + gap)
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(self.top),
                Inches(box_w),
                Inches(self.height),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = self.fill or theme.WASH
            box.line.color.rgb = theme.GREY
            box.text_frame.text = str(label)
            box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            _style_runs(box.text_frame, self.pt, theme.BODY_FONT, theme.INK)
            drawn.append(box)

            if self.arrows and i < count - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(x + box_w + 0.02),
                    Inches(self.top + self.height / 2 - 0.055),
                    Inches(gap - 0.04),
                    Inches(0.11),
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = theme.GREY
                arrow.line.fill.background()
                drawn.append(arrow)
        return drawn


@dataclass
class Code:
    text: str
    left: float
    top: float
    width: float
    height: float
    pt: float = 11.0

    @property
    def text_for_check(self):
        return self.text

    def draw(self, slide):
        box = slide.shapes.add_textbox(
            Inches(self.left),
            Inches(self.top),
            Inches(self.width),
            Inches(self.height),
        )
        frame = box.text_frame
        frame.word_wrap = False
        frame.text = self.text
        _style_runs(frame, self.pt, theme.MONO_FONT, theme.DARK)
        return box


@dataclass
class KeyValues:
    pairs: list
    left: float
    top: float
    width: float
    height: float
    pt: float = 12.0
    key_fraction: float = 0.42

    @property
    def text_for_check(self):
        return "\n".join(f"{k}\t{v}" for k, v in self.pairs)

    def draw(self, slide):
        key_w = self.width * self.key_fraction
        row_h = self.height / max(1, len(self.pairs))
        drawn = []
        for i, (key, value) in enumerate(self.pairs):
            y = self.top + i * row_h
            key_box = slide.shapes.add_textbox(
                Inches(self.left), Inches(y), Inches(key_w), Inches(row_h)
            )
            key_box.text_frame.text = str(key)
            _style_runs(key_box.text_frame, self.pt, theme.MONO_FONT, theme.MUTED)

            value_box = slide.shapes.add_textbox(
                Inches(self.left + key_w),
                Inches(y),
                Inches(self.width - key_w),
                Inches(row_h),
            )
            value_box.text_frame.text = str(value)
            _style_runs(value_box.text_frame, self.pt, theme.MONO_FONT, theme.INK)
            drawn.extend([key_box, value_box])
        return drawn

"""Slide descriptors and the renderer that turns them into pptx slides.

Content modules build Slide objects. Only this module knows python-pptx.
"""

from dataclasses import dataclass, field

from pptx import Presentation
from pptx.util import Inches, Pt

from deck import theme

LAYOUTS = (
    "TITLE",
    "SECTION_HEADER",
    "SECTION_HEADER_1",
    "TITLE_ONLY",
    "ONE_COLUMN_TEXT",
    "MAIN_POINT",
    "SECTION_TITLE_AND_DESCRIPTION",
    "CAPTION_ONLY",
    "BIG_NUMBER",
    "BLANK",
)

# Which Slide field feeds which placeholder index, per layout.
PLACEHOLDER_FIELD = {
    ("TITLE", 0): "title",
    ("TITLE", 1): "subtitle",
    ("SECTION_HEADER", 0): "title",
    ("SECTION_HEADER_1", 0): "title",
    ("TITLE_ONLY", 0): "title",
    ("ONE_COLUMN_TEXT", 0): "title",
    ("ONE_COLUMN_TEXT", 1): "body",
    ("MAIN_POINT", 0): "title",
    ("SECTION_TITLE_AND_DESCRIPTION", 0): "title",
    ("SECTION_TITLE_AND_DESCRIPTION", 1): "subtitle",
    ("SECTION_TITLE_AND_DESCRIPTION", 2): "body",
    ("CAPTION_ONLY", 1): "caption",
    ("BIG_NUMBER", 0): "big",
    ("BIG_NUMBER", 1): "body",
}


@dataclass
class Slide:
    layout: str
    title: str = ""
    subtitle: str = ""
    body: str = ""
    big: str = ""
    caption: str = ""
    shapes: list = field(default_factory=list)
    notes: str = ""
    seconds: int = 0
    pt_override: dict = field(default_factory=dict)


def open_template(path):
    """Open the template and remove the two slides it ships with."""
    prs = Presentation(str(path))
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)
    return prs


def _apply_text(placeholder, text, pt):
    frame = placeholder.text_frame
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(pt)
            run.font.name = theme.BODY_FONT


def render(prs, slide_spec):
    """Add one slide described by slide_spec and return the pptx slide."""
    if slide_spec.layout not in LAYOUTS:
        raise ValueError(f"unknown layout: {slide_spec.layout}")

    layouts = {layout.name: layout for layout in prs.slide_layouts}
    slide = prs.slides.add_slide(layouts[slide_spec.layout])

    for placeholder in list(slide.placeholders):
        idx = placeholder.placeholder_format.idx
        field_name = PLACEHOLDER_FIELD.get((slide_spec.layout, idx))
        text = getattr(slide_spec, field_name, "") if field_name else ""
        if text:
            pt = slide_spec.pt_override.get(
                idx, theme.PLACEHOLDER_PT[(slide_spec.layout, idx)]
            )
            _apply_text(placeholder, text, pt)
        else:
            placeholder._element.getparent().remove(placeholder._element)

    for shape_spec in slide_spec.shapes:
        shape_spec.draw(slide)

    slide.notes_slide.notes_text_frame.text = slide_spec.notes
    return slide


def build(template, section_slides):
    """Render every slide in order and return the Presentation."""
    prs = open_template(template)
    for slide_spec in section_slides:
        render(prs, slide_spec)
    return prs


def inches(*values):
    """Convenience for shape modules."""
    return tuple(Inches(v) for v in values)

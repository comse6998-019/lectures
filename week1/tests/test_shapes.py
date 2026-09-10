import pathlib

from pptx import Presentation

from deck import shapes, slides, theme

TEMPLATE = pathlib.Path(__file__).parent.parent.parent / "template.pptx"


def render_one(tmp_path, *shape_specs):
    prs = slides.build(
        TEMPLATE,
        [
            slides.Slide(
                layout="TITLE_ONLY",
                title="figure",
                shapes=list(shape_specs),
                notes="[60s] x",
                seconds=60,
            )
        ],
    )
    out = tmp_path / "fig.pptx"
    prs.save(out)
    return list(Presentation(out).slides)[0]


def test_table_renders_with_headers(tmp_path):
    table = shapes.Table(
        rows=[["run", "status"], ["run-smoke", "ok"]],
        left=0.4,
        top=0.9,
        width=6.7,
        height=1.2,
    )
    slide = render_one(tmp_path, table)
    rendered = [s for s in slide.shapes if s.has_table][0].table
    assert rendered.cell(0, 0).text == "run"
    assert rendered.cell(1, 1).text == "ok"


def test_table_exposes_text_for_redaction_check():
    table = shapes.Table(rows=[["a", "/Users/rkrsn"]], left=0, top=0, width=1, height=1)
    assert "/Users/rkrsn" in table.text_for_check


def test_boxes_draw_one_shape_per_label_plus_arrows(tmp_path):
    boxes = shapes.Boxes(
        labels=["model", "controller", "tools"],
        left=0.4,
        top=1.0,
        width=6.7,
        height=0.9,
    )
    slide = render_one(tmp_path, boxes)
    # 1 title placeholder + 3 boxes + 2 arrows
    assert len(slide.shapes) == 6


def test_code_block_uses_monospace(tmp_path):
    code = shapes.Code(
        text="run_start\nturn_start\nrun_end", left=0.4, top=0.9, width=6.7, height=2.0
    )
    slide = render_one(tmp_path, code)
    box = [s for s in slide.shapes if s.has_text_frame and "run_start" in s.text_frame.text][0]
    assert box.text_frame.paragraphs[0].runs[0].font.name == theme.MONO_FONT


def test_key_values_render_both_columns(tmp_path):
    kv = shapes.KeyValues(
        pairs=[("input", "57,944"), ("cache_read", "883,599,352")],
        left=0.4,
        top=0.9,
        width=6.7,
        height=1.6,
    )
    slide = render_one(tmp_path, kv)
    text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "cache_read" in text and "883,599,352" in text


def test_shapes_stay_inside_the_canvas():
    left, top, width, height = theme.CANVAS
    for spec in (
        shapes.Table(rows=[["a"]], left=left, top=top, width=width, height=1.0),
        shapes.Code(text="x", left=left, top=top, width=width, height=1.0),
    ):
        assert spec.left >= left
        assert spec.left + spec.width <= left + width + 1e-9
        assert spec.top >= top


def test_boxes_rendered_shapes_stay_inside_canvas(tmp_path):
    # Boxes derives each box's x from label count and gap arithmetic; check
    # the actual rendered geometry, not just the constructor's stored width.
    left, top, width, height = theme.CANVAS
    boxes = shapes.Boxes(
        labels=["model", "controller", "tools", "memory"],
        left=left,
        top=top,
        width=width,
        height=0.9,
    )
    slide = render_one(tmp_path, boxes)
    drawn = [s for s in slide.shapes if not s.is_placeholder]
    assert len(drawn) == 7  # 4 boxes + 3 arrows
    for s in drawn:
        assert s.left.inches >= left - 1e-6
        assert s.left.inches + s.width.inches <= left + width + 1e-6


def test_key_values_rendered_boxes_stay_inside_canvas(tmp_path):
    left, top, width, height = theme.CANVAS
    kv = shapes.KeyValues(
        pairs=[("input", "1"), ("output", "2"), ("cache_read", "3")],
        left=left,
        top=top,
        width=width,
        height=1.5,
    )
    slide = render_one(tmp_path, kv)
    drawn = [s for s in slide.shapes if not s.is_placeholder]
    assert len(drawn) == 6  # 3 key boxes + 3 value boxes
    for s in drawn:
        assert s.left.inches + s.width.inches <= left + width + 1e-6
        assert s.top.inches + s.height.inches <= top + height + 1e-6

import pathlib

import pytest
from pptx import Presentation

from deck import slides

TEMPLATE = pathlib.Path(__file__).parent.parent.parent / "template.pptx"


def test_template_is_stripped_of_seeded_slides():
    prs = slides.open_template(TEMPLATE)
    assert len(prs.slides._sldIdLst) == 0


def test_layout_names_are_the_templates_ten():
    prs = slides.open_template(TEMPLATE)
    assert set(slides.LAYOUTS) == {layout.name for layout in prs.slide_layouts}
    assert len(slides.LAYOUTS) == 10


def test_render_title_slide_roundtrips(tmp_path):
    prs = slides.open_template(TEMPLATE)
    slides.render(
        prs,
        slides.Slide(
            layout="TITLE",
            title="Design of Production Agentic Systems",
            subtitle="Lecture 1",
            notes="[180s] welcome",
            seconds=180,
        ),
    )
    out = tmp_path / "one.pptx"
    prs.save(out)

    reopened = Presentation(out)
    slide = list(reopened.slides)[0]
    assert slide.slide_layout.name == "TITLE"
    assert slide.placeholders[0].text_frame.text == (
        "Design of Production Agentic Systems"
    )
    assert slide.placeholders[1].text_frame.text == "Lecture 1"
    assert slide.notes_slide.notes_text_frame.text == "[180s] welcome"


def test_big_number_fills_title_and_label(tmp_path):
    prs = slides.open_template(TEMPLATE)
    slides.render(
        prs,
        slides.Slide(
            layout="BIG_NUMBER",
            big="28,972",
            body="model turns in one 9.1-hour run",
            notes="[45s] scale",
            seconds=45,
        ),
    )
    out = tmp_path / "big.pptx"
    prs.save(out)
    slide = list(Presentation(out).slides)[0]
    assert slide.placeholders[0].text_frame.text == "28,972"
    assert slide.placeholders[1].text_frame.text == (
        "model turns in one 9.1-hour run"
    )


def test_unfilled_placeholders_are_removed(tmp_path):
    prs = slides.open_template(TEMPLATE)
    slides.render(
        prs,
        slides.Slide(
            layout="SECTION_TITLE_AND_DESCRIPTION",
            title="Unit 1",
            subtitle="Who chooses the next operation?",
            notes="[30s] unit opener",
            seconds=30,
        ),
    )
    out = tmp_path / "partial.pptx"
    prs.save(out)
    slide = list(Presentation(out).slides)[0]
    # body (idx 2) was never filled, so it must not ship as an empty prompt box
    assert [p.placeholder_format.idx for p in slide.placeholders] == [0, 1]


def test_point_size_override_is_applied(tmp_path):
    prs = slides.open_template(TEMPLATE)
    slides.render(
        prs,
        slides.Slide(
            layout="BIG_NUMBER",
            big="883,599,352",
            body="cache_read tokens",
            pt_override={0: 54.0},
            notes="[60s] accounting",
            seconds=60,
        ),
    )
    out = tmp_path / "override.pptx"
    prs.save(out)
    slide = list(Presentation(out).slides)[0]
    run = slide.placeholders[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 54.0


def test_unknown_layout_is_rejected():
    prs = slides.open_template(TEMPLATE)
    with pytest.raises(ValueError, match="unknown layout"):
        slides.render(prs, slides.Slide(layout="TWO_COLUMN", notes="[1s] x"))
